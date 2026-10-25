#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script vận hành Skill mathplus-giai-toan:
Tạo tài liệu Word Lời giải chi tiết Đề kiểm tra giữa kì 1 Toán 10 THPT Trần Nhân Tông (Mã đề 294)
chuẩn nhận diện thương hiệu MathPlus Academy:
- Chữ Đen tuyền 100% (#000000), nền trắng hoàn toàn.
- Cỡ chữ toàn văn bản TỐI THIỂU 11pt.
- KHÔNG CÒN BẤT KỲ DẤU ** NÀO: Toàn bộ markdown **bold** được chuyển thành chữ in đậm (bold).
- ĐẦY ĐỦ HÌNH VẼ Ở CẢ PHẦN ĐỀ BÀI VÀ PHẦN LỜI GIẢI:
  + Đề bài Câu 4 (Phần I) có hình vẽ trục Oxy và đường thẳng d.
  + Đề bài Câu 1 (Phần III) có hình vẽ quan sát cây cao BC từ điểm A.
  + Đề bài Câu 2 (Phần IV) có hình vẽ tam giác ABC nội tiếp đường tròn.
  + Đề bài Câu 3 (Phần IV) có hình vẽ thiết bị chiếu sáng trên tường cao 5m và dải sáng.
  + Các câu bài toán quy hoạch (Câu 2, 3 Phần III; Câu 1 Phần IV) có đồ thị đa giác miền nghiệm trong phần Lời giải chi tiết.
- Khung lời giải chi tiết (Solution Box) KHÔNG có sọc đen bên trái.
- Chèn Logo MathPlus nguyên màu thương hiệu tại Header đầu trang.
- Header các trang sau chỉ gồm Logo và tên "Trung tâm MathPlus Academy" (11pt nghiêng).
- Chân trang đánh số trang tự động kèm Hotline 0977 961 189 và chữ ký Thầy Nguyễn Văn Dũng.
- 100% công thức toán học OMML (Word native).
"""

import os
import sys
import re

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.append(r"C:\Users\nguye\.gemini\config\skills\mathplus-giai-toan\scripts")
from mathplus_docx_engine import (
    MathPlusDocxBuilder,
    style_cell,
    add_rich_text,
    add_inline_math,
    add_display_math,
    COLOR_BLACK,
    Pt,
    Mm,
    Inches,
    RGBColor,
    WD_ALIGN_PARAGRAPH,
    WD_TABLE_ALIGNMENT
)
from omml_converter import latex_to_omml

def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_options_table(doc, options, cols=4):
    """Thêm bảng các phương án lựa chọn A, B, C, D không viền, căn lề thẳng hàng."""
    rows = 1 if cols == 4 else (2 if cols == 2 else 4)
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_w = Mm(172 / cols)
    for r in range(rows):
        for c in range(cols):
            idx = r * cols + c
            cell = tbl.cell(r, c)
            cell.width = col_w
            set_cell_margins(cell, 20, 20, 40, 40)
            style_cell(cell, bg_color=None, top="none", bottom="none", left="none", right="none")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if idx < len(options):
                add_rich_text(p, options[idx], default_size=Pt(11), default_color=COLOR_BLACK)
    return tbl

def add_question_figure(doc, fig_path, fig_width=Inches(3.2), caption_text=None):
    """Chèn hình vẽ trực tiếp vào phần câu hỏi của đề bài."""
    if os.path.exists(fig_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run()
        r.add_picture(fig_path, width=fig_width)
        if caption_text:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(0)
            p_cap.paragraph_format.space_after = Pt(3)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(11)
            r_cap.font.italic = True
            r_cap.font.color.rgb = COLOR_BLACK

def build_mathplus_solution_doc(output_path):
    print("Khởi tạo MathPlusDocxBuilder...")
    builder = MathPlusDocxBuilder(
        title="ĐỀ KIỂM TRA GIỮA HỌC KỲ I TOÁN 10",
        subtitle="Trường THPT Trần Nhân Tông – Mã đề 294 – Năm học 2026 – 2027",
        time_limit="90 phút",
        is_solved=True,
        logo_path=r"C:\code\mathplus\public\images\logo.jpg"
    )
    doc = builder.doc

    # =========================================================================
    # BẢNG TỔNG HỢP ĐÁP ÁN NHANH
    # =========================================================================
    builder.add_section(
        "BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ THI",
        "Tra cứu nhanh kết quả Phần I (Trắc nghiệm), Phần II (Đúng/Sai) và Phần III (Trả lời ngắn)"
    )

    # 1. Bảng Phần I
    p1_head = doc.add_paragraph()
    p1_head.paragraph_format.space_before = Pt(4)
    p1_head.paragraph_format.space_after = Pt(2)
    add_rich_text(p1_head, "**1. PHẦN I – Trắc nghiệm 4 lựa chọn (12 câu):**", default_size=Pt(11))

    t_ans1 = doc.add_table(rows=2, cols=12)
    t_ans1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_ans1.autofit = False
    ans1_keys = ["C", "A", "D", "A", "B", "C", "B", "A", "A", "A", "D", "A"]
    for c in range(12):
        c0 = t_ans1.cell(0, c); c1 = t_ans1.cell(1, c)
        c0.width = Mm(14.3); c1.width = Mm(14.3)
        style_cell(c0, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        style_cell(c1, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(2); p0.paragraph_format.space_after = Pt(2)
        add_rich_text(p0, f"**{c+1}**", default_size=Pt(11))
        p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(2); p1.paragraph_format.space_after = Pt(2)
        add_rich_text(p1, f"**{ans1_keys[c]}**", default_size=Pt(11))

    # 2. Bảng Phần II & Phần III đặt cạnh nhau
    p2_head = doc.add_paragraph()
    p2_head.paragraph_format.space_before = Pt(6)
    p2_head.paragraph_format.space_after = Pt(2)
    add_rich_text(p2_head, "**2. PHẦN II (Đúng/Sai) và PHẦN III (Trả lời ngắn):**", default_size=Pt(11))

    t_ans23 = doc.add_table(rows=3, cols=6)
    t_ans23.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_ans23.autofit = False
    h_data = ["Câu hỏi", "Ý a", "Ý b", "Ý c", "Ý d", "Phần III – Trả lời ngắn"]
    widths = [Mm(25), Mm(18), Mm(18), Mm(18), Mm(18), Mm(75)]
    for j, (h, w) in enumerate(zip(h_data, widths)):
        cell = t_ans23.cell(0, j); cell.width = w
        style_cell(cell, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, f"**{h}**", default_size=Pt(11))

    row1_data = ["Câu 1", "S", "Đ", "S", "Đ", "Câu 1: **14.5** (m)  |  Câu 2: **4500** (triệu đ)"]
    row2_data = ["Câu 2", "Đ", "Đ", "S", "Đ", "Câu 3: **72** (triệu đ)  |  Câu 4: **11.02**"]
    for i, rdata in enumerate([row1_data, row2_data]):
        for j, val in enumerate(rdata):
            cell = t_ans23.cell(i+1, j); cell.width = widths[j]
            style_cell(cell, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 5 else WD_ALIGN_PARAGRAPH.LEFT
            add_rich_text(p, f"**{val}**" if j in [1, 2, 3, 4] else val, default_size=Pt(11))

    # =========================================================================
    # PHẦN I. TRẮC NGHIỆM 4 LỰA CHỌN
    # =========================================================================
    builder.add_section(
        "PHẦN I. CÂU HỎI TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN",
        "Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi chỉ chọn một phương án chính xác."
    )

    # --- Câu 1 ---
    builder.add_problem_header("Câu 1.", "Bất phương trình nào sau đây **không** là bất phương trình bậc nhất hai ẩn?")
    add_options_table(doc, [
        "**A.** $x - 4y + 2 > 0.$",
        "**B.** $3x + 2y \\le 5.$",
        "**C.** $xy - 2x + 1 < 0.$",
        "**D.** $y + 7 \\ge 0.$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 1:'),
        ('p', r"Bất phương trình bậc nhất hai ẩn $x, y$ có dạng tổng quát là $ax + by + c < 0$ (hoặc $>0, \le 0, \ge 0$), trong đó $a, b, c$ là các số thực đã cho và $a, b$ không đồng thời bằng 0."),
        ('bullet', r"Phương án A: $x - 4y + 2 > 0$ có $a = 1, b = -4$ là BPT bậc nhất hai ẩn."),
        ('bullet', r"Phương án B: $3x + 2y \le 5 \iff 3x + 2y - 5 \le 0$ có $a = 3, b = 2$ là BPT bậc nhất hai ẩn."),
        ('bullet', r"Phương án C: $xy - 2x + 1 < 0$ chứa tích hai biến $xy$ (bậc hai đối với tập hợp biến), do đó **không** phải là bất phương trình bậc nhất hai ẩn."),
        ('bullet', r"Phương án D: $y + 7 \ge 0$ có $a = 0, b = 1$ vẫn là BPT bậc nhất hai ẩn (khuyết $x$)."),
        ('bullet', r"**Kết luận:** Chọn đáp án **C**.")
    ])

    # --- Câu 2 ---
    builder.add_problem_header("Câu 2.", "Bất phương trình $5x - 2(y - 2x - 3) > 0$ tương đương với bất phương trình nào sau đây?")
    add_options_table(doc, [
        "**A.** $9x - 2y + 6 > 0.$",
        "**B.** $x - 2y + 6 > 0.$",
        "**C.** $9x - 2y - 6 > 0.$",
        "**D.** $7x - 2y + 3 > 0.$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 2:'),
        ('p', r"Thực hiện khai triển và thu gọn vế trái của bất phương trình đã cho:"),
        ('display_math', r"5x - 2(y - 2x - 3) > 0 \iff 5x - 2y + 4x + 6 > 0 \iff 9x - 2y + 6 > 0."),
        ('bullet', r"**Kết luận:** Chọn đáp án **A**.")
    ])

    # --- Câu 3 ---
    builder.add_problem_header("Câu 3.", "Cặp số nào sau đây **không** là nghiệm của bất phương trình $2x + 3y - 6 > 0$?")
    add_options_table(doc, [
        "**A.** $(2; 1).$",
        "**B.** $(0; 3).$",
        "**C.** $(3; 1).$",
        "**D.** $(0; 0).$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 3:'),
        ('p', r"Đặt $f(x, y) = 2x + 3y - 6$. Cặp số $(x_0; y_0)$ là nghiệm nếu $f(x_0, y_0) > 0$:"),
        ('bullet', r"Tại $A(2; 1)$: $f(2, 1) = 2(2) + 3(1) - 6 = 1 > 0$ (thỏa mãn nghiệm)."),
        ('bullet', r"Tại $B(0; 3)$: $f(0, 3) = 2(0) + 3(3) - 6 = 3 > 0$ (thỏa mãn nghiệm)."),
        ('bullet', r"Tại $C(3; 1)$: $f(3, 1) = 2(3) + 3(1) - 6 = 3 > 0$ (thỏa mãn nghiệm)."),
        ('bullet', r"Tại $D(0; 0)$: $f(0, 0) = 2(0) + 3(0) - 6 = -6 < 0$ (không thỏa mãn $f(x, y) > 0$)."),
        ('bullet', r"**Kết luận:** Chọn đáp án **D**.")
    ])

    # --- Câu 4 (HÌNH VẼ Ở CẢ ĐỀ BÀI VÀ LỜI GIẢI) ---
    builder.add_problem_header("Câu 4.", "Miền không bị gạch chéo (không kể bờ $d$) trong hình dưới đây là miền nghiệm của bất phương trình nào?")
    # CHÈN HÌNH MINH HỌA NGAY TRONG ĐỀ BÀI
    add_question_figure(
        doc,
        r"C:\code\mathplus\lop10\figures_de294\c4_oxy.png",
        fig_width=Inches(3.3),
        caption_text="Hình vẽ: Miền nghiệm của bất phương trình trên mặt phẳng tọa độ Oxy"
    )
    add_options_table(doc, [
        "**A.** $x + 4y - 8 < 0.$",
        "**B.** $x + 4y - 8 > 0.$",
        "**C.** $4x + y - 8 < 0.$",
        "**D.** $4x + y - 8 > 0.$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 4:'),
        ('p', r"1. Xác định phương trình đường thẳng bờ $d$:"),
        ('bullet', r"Đường thẳng $d$ cắt trục tung $Oy$ tại điểm $A(0; 2)$ và cắt trục hoành $Ox$ tại điểm $B(8; 0)$."),
        ('bullet', r"Phương trình đường thẳng theo đoạn chắn: $\dfrac{x}{8} + \dfrac{y}{2} = 1 \iff x + 4y - 8 = 0$."),
        ('p', r"2. Xác định dấu của bất phương trình:"),
        ('bullet', r"Quan sát hình vẽ, miền không bị gạch chéo chứa gốc tọa độ $O(0; 0)$."),
        ('bullet', r"Thay tọa độ $O(0; 0)$ vào biểu thức $x + 4y - 8$: $0 + 4(0) - 8 = -8 < 0$."),
        ('bullet', r"Đề bài nêu rõ *không kể bờ $d$*, nên bất phương trình có dấu $<$ nghiêm ngặt: $x + 4y - 8 < 0$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **A**.")
    ])

    # --- Câu 5 ---
    builder.add_problem_header("Câu 5.", "Hệ bất phương trình nào sau đây **không** là hệ bất phương trình bậc nhất hai ẩn?")
    add_options_table(doc, [
        "**A.** $\\begin{cases} x + y \\le 5 \\\\ 2x - y > 0 \\end{cases}$",
        "**B.** $\\begin{cases} x^3 - y < 2 \\\\ x + 4y \\le 9 \\end{cases}$",
        "**C.** $\\begin{cases} x \\ge -1 \\\\ 3x + 2y < 7 \\end{cases}$",
        "**D.** $\\begin{cases} 4x - y \\ge 3 \\\\ y \\le 5 \\end{cases}$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 5:'),
        ('p', r"Hệ bất phương trình bậc nhất hai ẩn là hệ gồm hai hay nhiều bất phương trình, trong đó mỗi bất phương trình đều phải là bất phương trình bậc nhất hai ẩn."),
        ('bullet', r"Ở phương án B, bất phương trình $x^3 - y < 2$ có chứa lũy thừa bậc ba $x^3$, do đó không phải là bất phương trình bậc nhất hai ẩn."),
        ('bullet', r"**Kết luận:** Chọn đáp án **B**.")
    ])

    # --- Câu 6 ---
    builder.add_problem_header("Câu 6.", "Cặp số nào sau đây **không** là nghiệm của hệ bất phương trình $\\begin{cases} x + y - 3 \\le 0 \\\\ 2x - y + 1 > 0 \\end{cases}$?")
    add_options_table(doc, [
        "**A.** $(0; 0).$",
        "**B.** $(1; 1).$",
        "**C.** $(-1; 1).$",
        "**D.** $(2; 0).$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 6:'),
        ('p', r"Một cặp số $(x; y)$ là nghiệm của hệ khi và chỉ khi nó thỏa mãn đồng thời cả hai bất phương trình trong hệ:"),
        ('bullet', r"Với $A(0; 0)$: $0 + 0 - 3 = -3 \le 0$ (Đúng) và $2(0) - 0 + 1 = 1 > 0$ (Đúng)."),
        ('bullet', r"Với $B(1; 1)$: $1 + 1 - 3 = -1 \le 0$ (Đúng) và $2(1) - 1 + 1 = 2 > 0$ (Đúng)."),
        ('bullet', r"Với $C(-1; 1)$: Bất phương trình thứ hai trở thành $2(-1) - 1 + 1 = -2 > 0$ (SAI). Do đó $(-1; 1)$ không là nghiệm của hệ."),
        ('bullet', r"Với $D(2; 0)$: $2 + 0 - 3 = -1 \le 0$ (Đúng) và $2(2) - 0 + 1 = 5 > 0$ (Đúng)."),
        ('bullet', r"**Kết luận:** Chọn đáp án **C**.")
    ])

    # --- Câu 7 ---
    builder.add_problem_header("Câu 7.", "Cho hệ $\\begin{cases} x - \\dfrac{1}{4}y < 3 \\quad (1) \\\\ 4x - y < 8 \\quad (2) \\end{cases}$. Khẳng định nào sau đây đúng về tập nghiệm $S_1, S_2$ của $(1)$ và $(2)$?")
    add_options_table(doc, [
        "**A.** $S_1 \\subset S_2.$",
        "**B.** $S_2 \\subset S_1.$",
        "**C.** $S_1 = S_2.$",
        "**D.** $S_2 \\ne S_1.$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 7:'),
        ('p', r"Nhân cả hai vế của bất phương trình (1) với 4 (vì $4 > 0$ nên giữ nguyên chiều):"),
        ('display_math', r"x - \dfrac{1}{4}y < 3 \iff 4x - y < 12."),
        ('bullet', r"Tập nghiệm $S_1 = \{(x, y) \in \mathbb{R}^2 \mid 4x - y < 12\}$."),
        ('bullet', r"Tập nghiệm $S_2 = \{(x, y) \in \mathbb{R}^2 \mid 4x - y < 8\}$."),
        ('bullet', r"Với mọi cặp số $(x, y) \in S_2$, ta luôn có $4x - y < 8$. Vì $8 < 12$ nên hiển nhiên suy ra $4x - y < 12$, tức là $(x, y) \in S_1$."),
        ('bullet', r"Do đó, mọi phần tử của $S_2$ đều thuộc $S_1$, suy ra $S_2 \subset S_1$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **B**.")
    ])

    # --- Câu 8 ---
    builder.add_problem_header("Câu 8.", "Cho $\\alpha$ là góc tù và $\\sin\\alpha = \\dfrac{9}{41}$. Tính $\\cos\\alpha$.")
    add_options_table(doc, [
        "**A.** $-\\dfrac{40}{41}.$",
        "**B.** $\\dfrac{40}{41}.$",
        "**C.** $-\\dfrac{32}{41}.$",
        "**D.** $\\dfrac{32}{41}.$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 8:'),
        ('bullet', r"Vì $\alpha$ là góc tù ($90^\circ < \alpha < 180^\circ$) nên điểm ngọn của góc thuộc góc phần tư thứ II, do đó $\cos\alpha < 0$."),
        ('bullet', r"Áp dụng hệ thức lượng giác cơ bản $\sin^2\alpha + \cos^2\alpha = 1$:"),
        ('display_math', r"\cos^2\alpha = 1 - \sin^2\alpha = 1 - \left(\dfrac{9}{41}\right)^2 = 1 - \dfrac{81}{1681} = \dfrac{1600}{1681} = \left(\dfrac{40}{41}\right)^2."),
        ('bullet', r"Vì $\cos\alpha < 0$ nên ta chọn nghiệm âm: $\cos\alpha = -\sqrt{\dfrac{1600}{1681}} = -\dfrac{40}{41}$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **A**.")
    ])

    # --- Câu 9 ---
    builder.add_problem_header("Câu 9.", "Cho $\\sin\\alpha = \\dfrac{20}{29}$ với $90^\\circ < \\alpha < 180^\\circ$. Tính $\\tan(180^\\circ - \\alpha)$.")
    add_options_table(doc, [
        "**A.** $\\dfrac{20}{21}.$",
        "**B.** $-\\dfrac{20}{21}.$",
        "**C.** $\\dfrac{21}{20}.$",
        "**D.** $-\\dfrac{21}{20}.$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 9:'),
        ('bullet', r"Do $90^\circ < \alpha < 180^\circ \implies \cos\alpha < 0$."),
        ('bullet', r"Ta có: $\cos^2\alpha = 1 - \sin^2\alpha = 1 - \left(\dfrac{20}{29}\right)^2 = 1 - \dfrac{400}{841} = \dfrac{441}{841} = \left(\dfrac{21}{29}\right)^2$."),
        ('bullet', r"Suy ra $\cos\alpha = -\dfrac{21}{29}$."),
        ('bullet', r"Tính giá trị $\tan\alpha$: $\tan\alpha = \dfrac{\sin\alpha}{\cos\alpha} = \dfrac{20/29}{-21/29} = -\dfrac{20}{21}$."),
        ('bullet', r"Áp dụng tính chất hai góc bù nhau: $\tan(180^\circ - \alpha) = -\tan\alpha = -\left(-\dfrac{20}{21}\right) = \dfrac{20}{21}$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **A**.")
    ])

    # --- Câu 10 ---
    builder.add_problem_header("Câu 10.", "Cho tam giác $ABC$ có $b = 8, c = 6$ và $\\cos A = \\dfrac{3}{4}$. Tính độ dài cạnh $a$.")
    add_options_table(doc, [
        "**A.** $2\\sqrt{7}.$",
        "**B.** $4.$",
        "**C.** $\\sqrt{22}.$",
        "**D.** $6.$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 10:'),
        ('p', r"Áp dụng định lý côsin trong tam giác $ABC$:"),
        ('display_math', r"a^2 = b^2 + c^2 - 2bc\cos A"),
        ('bullet', r"Thay các số liệu đã cho vào công thức:"),
        ('display_math', r"a^2 = 8^2 + 6^2 - 2 \cdot 8 \cdot 6 \cdot \dfrac{3}{4} = 64 + 36 - 72 = 28."),
        ('bullet', r"Vì độ dài cạnh $a > 0$ nên $a = \sqrt{28} = \sqrt{4 \cdot 7} = 2\sqrt{7}$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **A**.")
    ])

    # --- Câu 11 ---
    builder.add_problem_header("Câu 11.", "Cho tam giác $ABC$ và bán kính đường tròn ngoại tiếp $R$. Khẳng định nào sau đây sai?")
    add_options_table(doc, [
        "**A.** $\\sin A = \\dfrac{a}{2R}.$",
        "**B.** $\\dfrac{c}{\\sin C} = 2R.$",
        "**C.** $a\\sin B = b\\sin A.$",
        "**D.** $c\\sin C = 2R.$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 11:'),
        ('p', r"Định lý sin trong tam giác phát biểu: $\dfrac{a}{\sin A} = \dfrac{b}{\sin B} = \dfrac{c}{\sin C} = 2R$."),
        ('bullet', r"Phương án A: $\dfrac{a}{\sin A} = 2R \iff \sin A = \dfrac{a}{2R}$ (Đúng)."),
        ('bullet', r"Phương án B: $\dfrac{c}{\sin C} = 2R$ (Đúng)."),
        ('bullet', r"Phương án C: $\dfrac{a}{\sin A} = \dfrac{b}{\sin B} \iff a\sin B = b\sin A$ (Đúng)."),
        ('bullet', r"Phương án D: $c\sin C = 2R$ là Sai vì theo định lý sin phải là $\dfrac{c}{\sin C} = 2R$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **D**.")
    ])

    # --- Câu 12 ---
    builder.add_problem_header("Câu 12.", "Cho tam giác $ABC$ có $a = 10, c = 6$ và $\\widehat{B} = 30^\\circ$. Diện tích của tam giác $ABC$ bằng bao nhiêu?")
    add_options_table(doc, [
        "**A.** $15.$",
        "**B.** $30.$",
        "**C.** $15\\sqrt{3}.$",
        "**D.** $10\\sqrt{3}.$"
    ], cols=4)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 12:'),
        ('p', r"Diện tích tam giác được tính theo công thức nửa tích hai cạnh nhân sin của góc xen giữa:"),
        ('display_math', r"S = \dfrac{1}{2}ac\sin B"),
        ('bullet', r"Thay số: $S = \dfrac{1}{2} \cdot 10 \cdot 6 \cdot \sin 30^\circ = 30 \cdot \dfrac{1}{2} = 15$."),
        ('bullet', r"**Kết luận:** Chọn đáp án **A**.")
    ])

    # =========================================================================
    # PHẦN II. TRẮC NGHIỆM ĐÚNG SAI
    # =========================================================================
    builder.add_section(
        "PHẦN II. CÂU HỎI TRẮC NGHIỆM ĐÚNG SAI",
        "Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn Đúng hoặc Sai."
    )

    # --- Câu 1 Phần II ---
    builder.add_problem_header("Câu 1.", "Cho tam giác $ABC$ có $\\widehat{A} = 60^\\circ, AB = 9, AC = 15$. Xét tính đúng/sai của các mệnh đề sau:")
    add_options_table(doc, [
        "**a)** $S = \\dfrac{1}{2}AB \\cdot AC \\cos A.$",
        "**b)** $S = \\dfrac{135\\sqrt{3}}{4}.$",
        "**c)** $h_b = 9\\sqrt{3}.$",
        "**d)** $R = \\sqrt{57}.$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 1 (PHẦN II):'),
        ('p', r"Xét lần lượt từng mệnh đề toán học:"),
        ('bullet', r"**Mệnh đề a)** $S = \dfrac{1}{2}AB \cdot AC \cos A$: **SAI**."),
        ('p', r"Giải thích: Công thức diện tích tam giác theo hai cạnh và góc xen giữa là $S = \dfrac{1}{2}bc\sin A = \dfrac{1}{2}AB \cdot AC \sin A$. Do đó mệnh đề a là SAI."),
        ('bullet', r"**Mệnh đề b)** $S = \dfrac{135\sqrt{3}}{4}$: **ĐÚNG**."),
        ('p', r"Tính toán: $S = \dfrac{1}{2} \cdot 9 \cdot 15 \cdot \sin 60^\circ = \dfrac{135}{2} \cdot \dfrac{\sqrt{3}}{2} = \dfrac{135\sqrt{3}}{4}$."),
        ('bullet', r"**Mệnh đề c)** $h_b = 9\sqrt{3}$: **SAI**."),
        ('p', r"Giải thích: Chiều cao hạ từ đỉnh B xuống cạnh AC ($b = 15$) là: $S = \dfrac{1}{2}b \cdot h_b \iff h_b = \dfrac{2S}{b} = \dfrac{2 \cdot \frac{135\sqrt{3}}{4}}{15} = \dfrac{135\sqrt{3}}{30} = \dfrac{9\sqrt{3}}{2} \ne 9\sqrt{3}$."),
        ('bullet', r"**Mệnh đề d)** $R = \sqrt{57}$: **ĐÚNG**."),
        ('p', r"Tính toán:"),
        ('display_math', r"a^2 = b^2 + c^2 - 2bc\cos A = 15^2 + 9^2 - 2 \cdot 15 \cdot 9 \cdot \cos 60^\circ = 225 + 81 - 135 = 171 \implies a = \sqrt{171} = 3\sqrt{19}."),
        ('p', r"Theo định lý sin: $R = \dfrac{a}{2\sin A} = \dfrac{3\sqrt{19}}{2\sin 60^\circ} = \dfrac{3\sqrt{19}}{2 \cdot \frac{\sqrt{3}}{2}} = \dfrac{3\sqrt{19}}{\sqrt{3}} = \sqrt{3 \cdot 19} = \sqrt{57}$."),
        ('bullet', r"**Kết luận:** a) **Sai**, b) **Đúng**, c) **Sai**, d) **Đúng**.")
    ])

    # --- Câu 2 Phần II ---
    builder.add_problem_header("Câu 2.", "Cho tam giác $ABC$ có độ dài ba cạnh lần lượt là $5, 12, 13$. Xét tính đúng/sai của các mệnh đề sau:")
    add_options_table(doc, [
        "**a)** $S = 30.$",
        "**b)** $p = 15.$",
        "**c)** $r = 3.$",
        "**d)** $R = \\dfrac{13}{2}.$"
    ], cols=2)
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 2 (PHẦN II):'),
        ('p', r"Nhận xét đặc biệt: Ta có $5^2 + 12^2 = 25 + 144 = 169 = 13^2$. Theo định lý Pytago đảo, $\triangle ABC$ vuông tại góc đối diện với cạnh huyền có độ dài bằng 13."),
        ('bullet', r"**Mệnh đề a)** $S = 30$: **ĐÚNG**."),
        ('p', r"Diện tích tam giác vuông bằng nửa tích hai cạnh góc vuông: $S = \dfrac{1}{2} \cdot 5 \cdot 12 = 30$."),
        ('bullet', r"**Mệnh đề b)** $p = 15$: **ĐÚNG**."),
        ('p', r"Nửa chu vi tam giác: $p = \dfrac{a + b + c}{2} = \dfrac{5 + 12 + 13}{2} = \dfrac{30}{2} = 15$."),
        ('bullet', r"**Mệnh đề c)** $r = 3$: **SAI**."),
        ('p', r"Bán kính đường tròn nội tiếp: $r = \dfrac{S}{p} = \dfrac{30}{15} = 2 \ne 3$. Do đó mệnh đề c là SAI."),
        ('bullet', r"**Mệnh đề d)** $R = \dfrac{13}{2}$: **ĐÚNG**."),
        ('p', r"Trong tam giác vuông, tâm đường tròn ngoại tiếp là trung điểm cạnh huyền và bán kính bằng nửa cạnh huyền: $R = \dfrac{13}{2} = 6.5$."),
        ('bullet', r"**Kết luận:** a) **Đúng**, b) **Đúng**, c) **Sai**, d) **Đúng**.")
    ])

    # =========================================================================
    # PHẦN III. TRẢ LỜI NGẮN
    # =========================================================================
    builder.add_section(
        "PHẦN III. CÂU HỎI TRẢ LỜI NGẮN",
        "Thí sinh trả lời từ câu 1 đến câu 4. Viết kết quả vào ô trả lời tương ứng."
    )

    # --- Câu 1 Phần III (HÌNH VẼ Ở CẢ ĐỀ BÀI VÀ LỜI GIẢI) ---
    builder.add_problem_header(
        "Câu 1.",
        r"Từ vị trí $A$, người ta quan sát thấy một cây cao, biết $AH = 4\text{ m}, HB = 15\text{ m}, \widehat{BAC} = 50^\circ$. "
        r"Khi đó chiều cao của cây là bao nhiêu? *(Lấy kết quả làm tròn một chữ số thập phân sau dấu phẩy)*."
    )
    # CHÈN HÌNH MINH HỌA NGAY TRONG ĐỀ BÀI
    add_question_figure(
        doc,
        r"C:\code\mathplus\lop10\figures_de294\p3_c1_tree.png",
        fig_width=Inches(3.3),
        caption_text="Hình vẽ: Mô hình quan sát cây cao BC từ vị trí quan sát A"
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 1 (PHẦN III):'),
        ('p', r"1. Phân tích mô hình hình học:"),
        ('bullet', r"Cây mọc vuông góc với mặt đất $HB$ tại điểm $B$, do đó $BC \perp HB$."),
        ('bullet', r"Vị trí quan sát $A$ có hình chiếu vuông góc lên mặt đất là $H$, do đó $AH \perp HB$ tại $H$."),
        ('bullet', r"Kẻ đường nằm ngang $AK \perp BC$ tại $K$. Tứ giác $AHBK$ có 3 góc vuông ($\widehat{H} = \widehat{B} = \widehat{K} = 90^\circ$) nên $AHBK$ là hình chữ nhật."),
        ('bullet', r"Suy ra: $AK = HB = 15\text{ m}$ và đoạn dưới của cây $BK = AH = 4\text{ m}$."),
        ('p', r"2. Tính toán các góc và chiều cao đoạn trên $KC$:"),
        ('bullet', r"Trong tam giác vuông $AKB$ tại $K$:"),
        ('display_math', r"\tan(\widehat{KAB}) = \dfrac{BK}{AK} = \dfrac{4}{15} \implies \widehat{KAB} = \arctan\left(\dfrac{4}{15}\right) \approx 14.931^\circ."),
        ('bullet', r"Do tia $AB$ nằm giữa hai tia $AK$ và $AC$, ta có góc nâng từ đường nằm ngang đến ngọn cây:"),
        ('display_math', r"\widehat{KAC} = \widehat{BAC} - \widehat{KAB} \approx 50^\circ - 14.931^\circ = 35.069^\circ."),
        ('bullet', r"Trong tam giác vuông $AKC$ tại $K$:"),
        ('display_math', r"KC = AK \cdot \tan(\widehat{KAC}) = 15 \cdot \tan(35.069^\circ) \approx 15 \cdot 0.70196 \approx 10.529\text{ m}."),
        ('p', r"3. Chiều cao toàn bộ của cây:"),
        ('display_math', r"h = BC = BK + KC = 4 + 10.529 = 14.529\text{ m} \approx 14.5\text{ m}."),
        ('bullet', r"**Đáp số điền ô:** **14.5**")
    ])

    # --- Câu 2 Phần III (KÈM ĐỒ THỊ ĐA GIÁC TRONG LỜI GIẢI) ---
    builder.add_problem_header(
        "Câu 2.",
        "Một cửa hàng dự định kinh doanh bổ sung hai dòng sản phẩm dòng M và dòng N, với số vốn ban đầu không vượt quá 39 tỉ đồng. "
        "Giá nhập một sản phẩm dòng M là 50 triệu đồng, lợi nhuận dự kiến là 6 triệu đồng; giá nhập một sản phẩm dòng N là 30 triệu đồng, lợi nhuận dự kiến là 3 triệu đồng. "
        "Tổng nhu cầu thị trường không vượt quá 1000 sản phẩm và nhu cầu dòng M không vượt quá 2 lần nhu cầu dòng N. "
        "Hỏi lợi nhuận lớn nhất có thể thu được là bao nhiêu *(đơn vị: triệu đồng)*?"
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 2 (PHẦN III):'),
        ('p', r"1. Thiết lập hệ bất phương trình ràng buộc:"),
        ('bullet', r"Gọi $x, y$ lần lượt là số lượng sản phẩm dòng M và dòng N cửa hàng nhập để bán ($x, y \ge 0, x, y \in \mathbb{N}$)."),
        ('bullet', r"Ràng buộc về nguồn vốn: $50x + 30y \le 39\,000 \iff 5x + 3y \le 3\,900$ (triệu đồng)."),
        ('bullet', r"Ràng buộc về tổng nhu cầu thị trường: $x + y \le 1\,000$."),
        ('bullet', r"Ràng buộc về tương quan nhu cầu hai dòng: $x \le 2y$."),
        ('p', r"Hàm mục tiêu lợi nhuận cần cực đại hóa: $F(x, y) = 6x + 3y$ (triệu đồng)."),
        ('p', r"2. Xác định các đỉnh của miền nghiệm đa giác:"),
        ('bullet', r"Gốc tọa độ $O(0; 0) \implies F(0, 0) = 0$ triệu đồng."),
        ('bullet', r"Điểm $A(0; 1000)$: $F(0, 1000) = 6(0) + 3(1000) = 3\,000\text{ triệu đồng}$."),
        ('bullet', r"Điểm $B$ là giao điểm của $x + y = 1000$ và $5x + 3y = 3900$: $B(450; 550) \implies F(450, 550) = 6(450) + 3(550) = 4\,350\text{ triệu đồng}$."),
        ('bullet', r"Điểm $C$ là giao điểm của $x = 2y$ và $5x + 3y = 3900$: thay $x = 2y \implies 13y = 3900 \iff y = 300, x = 600$."),
        ('p', r"Kiểm tra điều kiện: $x + y = 600 + 300 = 900 \le 1000$ (thỏa mãn!)."),
        ('display_math', r"F(600, 300) = 6(600) + 3(300) = 3600 + 900 = 4\,500\text{ triệu đồng}."),
        ('image', r"C:\code\mathplus\lop10\figures_de294\p3_c2_polygon.png", Inches(3.2)),
        ('caption', r"Hình vẽ: Miền nghiệm bài toán kinh doanh và điểm cực đại C(600; 300)"),
        ('p', r"3. So sánh và kết luận:"),
        ('bullet', r"Ta thấy $4\,500 > 4\,350 > 3\,000 > 0$. Vậy lợi nhuận tối đa thu được là 4500 triệu đồng khi kinh doanh 600 sản phẩm M và 300 sản phẩm N."),
        ('bullet', r"**Đáp số điền ô:** **4500**")
    ])

    # --- Câu 3 Phần III (KÈM ĐỒ THỊ ĐA GIÁC TRONG LỜI GIẢI) ---
    builder.add_problem_header(
        "Câu 3.",
        r"Một hộ nông dân có $1\,200\text{ m}^2$ đất. Trên mỗi $100\text{ m}^2$, loại cây thứ nhất cần 20 công và thu 5 triệu đồng; "
        r"loại thứ hai cần 30 công và thu 7 triệu đồng. Tổng số công không quá 300. Hỏi thu nhập lớn nhất là bao nhiêu triệu đồng?"
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 3 (PHẦN III):'),
        ('p', r"1. Gọi ẩn và đặt điều kiện:"),
        ('bullet', r"Gọi $x, y$ lần lượt là diện tích đất dùng để trồng loại cây thứ nhất và thứ hai (đơn vị: $100\text{ m}^2$, điều kiện $x, y \ge 0$)."),
        ('bullet', r"Giới hạn tổng diện tích: $100x + 100y \le 1200 \iff x + y \le 12$."),
        ('bullet', r"Giới hạn tổng số ngày công: $20x + 30y \le 300 \iff 2x + 3y \le 30$."),
        ('p', r"2. Hàm mục tiêu thu nhập: $T(x, y) = 5x + 7y$ (triệu đồng)."),
        ('p', r"3. Tìm giá trị tại các đỉnh miền nghiệm đa giác:"),
        ('bullet', r"Đỉnh $O(0; 0) \implies T(0, 0) = 0$ triệu đồng."),
        ('bullet', r"Đỉnh $A(12; 0) \implies T(12, 0) = 5(12) = 60$ triệu đồng."),
        ('bullet', r"Đỉnh $B(0; 10) \implies T(0, 10) = 7(10) = 70$ triệu đồng."),
        ('bullet', r"Đỉnh $C$ là giao điểm của $x + y = 12$ và $2x + 3y = 30$: giải ra $x = 6, y = 6$."),
        ('p', r"Giá trị thu nhập tại $C(6; 6)$: $T(6, 6) = 5(6) + 7(6) = 30 + 42 = 72$ triệu đồng."),
        ('image', r"C:\code\mathplus\lop10\figures_de294\p3_c3_polygon.png", Inches(3.2)),
        ('caption', r"Hình vẽ: Miền nghiệm quy hoạch đất trồng và điểm tối ưu C(6; 6)"),
        ('bullet', r"So sánh: $72 > 70 > 60 > 0$. Vậy thu nhập lớn nhất hộ nông dân đạt được là 72 triệu đồng."),
        ('bullet', r"**Đáp số điền ô:** **72**")
    ])

    # --- Câu 4 Phần III ---
    builder.add_problem_header(
        "Câu 4.",
        r"Cho tam giác $ABC$ có $BC = 9, \widehat{B} = 60^\circ, \widehat{C} = 75^\circ$. "
        r"Tính độ dài cạnh $AC$ *(làm tròn đến hàng phần trăm)*."
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 4 (PHẦN III):'),
        ('bullet', r"Tính góc $\widehat{A}$ trong tam giác $ABC$:"),
        ('display_math', r"\widehat{A} = 180^\circ - (\widehat{B} + \widehat{C}) = 180^\circ - (60^\circ + 75^\circ) = 180^\circ - 135^\circ = 45^\circ."),
        ('bullet', r"Áp dụng định lý sin trong tam giác $ABC$:"),
        ('display_math', r"\dfrac{AC}{\sin B} = \dfrac{BC}{\sin A} \implies AC = BC \cdot \dfrac{\sin B}{\sin A}."),
        ('bullet', r"Thay số cụ thể:"),
        ('display_math', r"AC = 9 \cdot \dfrac{\sin 60^\circ}{\sin 45^\circ} = 9 \cdot \dfrac{\frac{\sqrt{3}}{2}}{\frac{\sqrt{2}}{2}} = 9 \cdot \dfrac{\sqrt{3}}{\sqrt{2}} = \dfrac{9\sqrt{6}}{2}."),
        ('bullet', r"Tính giá trị thập phân: $\dfrac{9\sqrt{6}}{2} \approx 11.0227\dots \approx 11.02$."),
        ('bullet', r"**Đáp số điền ô:** **11.02**")
    ])

    # =========================================================================
    # PHẦN IV. TỰ LUẬN
    # =========================================================================
    builder.add_section(
        "PHẦN IV. CÂU HỎI TỰ LUẬN BÀI BẢN",
        "Thí sinh trình bày chi tiết các bước lập luận, biến đổi hình học và kết luận đầy đủ."
    )

    # --- Câu 1 Phần IV (KÈM ĐỒ THỊ ĐA GIÁC TRONG LỜI GIẢI) ---
    builder.add_problem_header(
        "Câu 1.",
        "Một cơ sở sản xuất sử dụng hai loại nguyên liệu $K$ và $P$. Mỗi tấn nguyên liệu $K$ có giá 4 triệu đồng, "
        "chiết xuất được 30 kg sản phẩm I và 30 kg sản phẩm II; mỗi tấn nguyên liệu $P$ có giá 5 triệu đồng, "
        "chiết xuất được 20 kg sản phẩm I và 40 kg sản phẩm II. "
        "Cơ sở cần sản xuất ít nhất 90 kg sản phẩm I và 120 kg sản phẩm II, đồng thời kho chứa chỉ cho phép mua tối đa không quá 4 tấn mỗi loại. "
        "Hãy tìm phương án mua nguyên liệu sao cho tổng chi phí là nhỏ nhất."
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 1 (PHẦN IV):'),
        ('p', r"**Bước 1: Gọi ẩn và đặt điều kiện thực tế**"),
        ('bullet', r"Gọi $x$ là số tấn nguyên liệu loại $K$ cần mua ($0 \le x \le 4$)."),
        ('bullet', r"Gọi $y$ là số tấn nguyên liệu loại $P$ cần mua ($0 \le y \le 4$)."),
        ('p', r"**Bước 2: Lập hệ bất phương trình theo yêu cầu sản phẩm**"),
        ('bullet', r"Sản phẩm loại I: $30x + 20y \ge 90 \iff 3x + 2y \ge 9$."),
        ('bullet', r"Sản phẩm loại II: $30x + 40y \ge 120 \iff 3x + 4y \ge 12$."),
        ('bullet', r"Hệ bất phương trình ràng buộc: $\begin{cases} 3x + 2y \ge 9 \\ 3x + 4y \ge 12 \\ 0 \le x \le 4 \\ 0 \le y \le 4 \end{cases}$"),
        ('p', r"**Bước 3: Khảo sát chi phí tại các đỉnh miền nghiệm**"),
        ('p', r"Hàm chi phí: $C(x, y) = 4x + 5y$ (triệu đồng)."),
        ('bullet', r"Đỉnh $M(1/3; 4)$: $C(M) = 4(1/3) + 5(4) = \dfrac{64}{3} \approx 21.33\text{ triệu đồng}$."),
        ('bullet', r"Đỉnh $N(2; 1.5)$ (giao điểm của $3x + 2y = 9$ và $3x + 4y = 12$): $C(N) = 4(2) + 5(1.5) = 15.5\text{ triệu đồng}$."),
        ('bullet', r"Đỉnh $P(4; 0)$: $C(P) = 4(4) + 5(0) = 16\text{ triệu đồng}$."),
        ('bullet', r"Đỉnh $Q(4; 4)$: $C(Q) = 4(4) + 5(4) = 36\text{ triệu đồng}$."),
        ('image', r"C:\code\mathplus\lop10\figures_de294\p4_c1_polygon.png", Inches(3.2)),
        ('caption', r"Hình vẽ: Đa giác miền nghiệm bài toán chi phí và điểm cực tiểu N(2; 1.5)"),
        ('p', r"**Bước 4: Kết luận phương án tối ưu**"),
        ('bullet', r"So sánh: $15.5 < 16 < 21.33 < 36$."),
        ('bullet', r"**Kết luận:** Cơ sở cần mua **2 tấn nguyên liệu K** và **1.5 tấn nguyên liệu P** để đạt chi phí thấp nhất là **15.5 triệu đồng** (15 triệu 500 nghìn đồng).")
    ])

    # --- Câu 2 Phần IV (HÌNH VẼ Ở CẢ ĐỀ BÀI VÀ LỜI GIẢI) ---
    builder.add_problem_header(
        "Câu 2.",
        "Từ một tấm bìa hình tròn, người ta cắt ra một tam giác $ABC$ có $AB = 5\\text{ cm}, AC = 7\\text{ cm}$ và $\\widehat{B} = 60^\\circ$. "
        "Hãy tính độ dài cạnh $BC$ và bán kính $R$ của tấm bìa hình tròn ban đầu *(kết quả làm tròn $R$ đến hàng phần mười)*."
    )
    # CHÈN HÌNH MINH HỌA NGAY TRONG ĐỀ BÀI
    add_question_figure(
        doc,
        r"C:\code\mathplus\lop10\figures_de294\p4_c2_circle.png",
        fig_width=Inches(2.5),
        caption_text="Hình vẽ: Tam giác ABC nội tiếp trong tấm bìa hình tròn"
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 2 (PHẦN IV):'),
        ('p', r"**1. Tính độ dài cạnh $BC$:**"),
        ('bullet', r"Áp dụng định lý côsin trong $\triangle ABC$ cho góc $\widehat{B} = 60^\circ$:"),
        ('display_math', r"AC^2 = AB^2 + BC^2 - 2AB \cdot BC \cos B"),
        ('bullet', r"Thay độ dài $AB = 5\text{ cm}, AC = 7\text{ cm}, \cos 60^\circ = \dfrac{1}{2}$:"),
        ('display_math', r"7^2 = 5^2 + BC^2 - 2 \cdot 5 \cdot BC \cdot \dfrac{1}{2} \iff 49 = 25 + BC^2 - 5BC."),
        ('display_math', r"BC^2 - 5BC - 24 = 0 \iff (BC - 8)(BC + 3) = 0."),
        ('bullet', r"Vì độ dài đoạn thẳng $BC > 0$ nên ta nhận $BC = 8\text{ cm}$."),
        ('p', r"**2. Tính bán kính $R$ của hình tròn:**"),
        ('bullet', r"Tấm bìa hình tròn ban đầu chính là hình tròn ngoại tiếp tam giác $ABC$."),
        ('bullet', r"Áp dụng định lý sin:"),
        ('display_math', r"\dfrac{AC}{\sin B} = 2R \implies R = \dfrac{AC}{2\sin B} = \dfrac{7}{2\sin 60^\circ} = \dfrac{7}{2 \cdot \frac{\sqrt{3}}{2}} = \dfrac{7}{\sqrt{3}} = \dfrac{7\sqrt{3}}{3} \approx 4.0415\text{ cm} \approx 4.0\text{ cm}."),
        ('bullet', r"**Kết luận:** Độ dài cạnh $BC = 8\text{ cm}$ và bán kính đường tròn $R \approx 4.0\text{ cm}$.")
    ])

    # --- Câu 3 Phần IV (HÌNH VẼ Ở CẢ ĐỀ BÀI VÀ LỜI GIẢI) ---
    builder.add_problem_header(
        "Câu 3.",
        "Người ta lắp một thiết bị chiếu sáng sát tường tại điểm $A$, cao $5\\text{ m}$ so với mặt đất. "
        "Góc chiếu sáng của thiết bị là $30^\\circ$ và dải sáng bắt đầu rọi xuống mặt đất tại điểm cách chân tường $3\\text{ m}$. "
        "Hỏi độ dài vùng được chiếu sáng trên mặt đất bằng bao nhiêu mét? *(Kết quả làm tròn đến hàng phần mười)*."
    )
    # CHÈN HÌNH MINH HỌA NGAY TRONG ĐỀ BÀI
    add_question_figure(
        doc,
        r"C:\code\mathplus\lop10\figures_de294\p4_c3_light.png",
        fig_width=Inches(3.2),
        caption_text="Hình vẽ: Thiết bị chiếu sáng tại điểm A sát tường và dải sáng trên mặt đất"
    )
    builder.add_solution([
        ('title', '✎ LỜI GIẢI CHI TIẾT CÂU 3 (PHẦN IV):'),
        ('p', r"**1. Mô hình hóa hình học bài toán:**"),
        ('bullet', r"Bức tường thẳng đứng $AB = 5\text{ m}$ vuông góc với mặt đất $DB$ tại điểm $B$ ($\widehat{B} = 90^\circ$)."),
        ('bullet', r"Điểm bắt đầu dải sáng trên mặt đất là $C$, cách chân tường $B$ một đoạn $BC = 3\text{ m}$."),
        ('bullet', r"Điểm kết thúc dải sáng trên mặt đất là $D$. Tia sáng ngoài cùng là $AD$."),
        ('bullet', r"Góc chiếu sáng của thiết bị phát ra là góc $\widehat{CAD} = 30^\circ$."),
        ('bullet', r"Độ dài vùng được chiếu sáng trên mặt đất chính là chiều dài đoạn thẳng $CD$."),
        ('p', r"**2. Tính toán lượng giác:**"),
        ('bullet', r"Trong tam giác vuông $ABC$ vuông tại $B$:"),
        ('display_math', r"\tan(\widehat{BAC}) = \dfrac{BC}{AB} = \dfrac{3}{5} = 0.6 \implies \widehat{BAC} = \arctan(0.6) \approx 30.964^\circ."),
        ('bullet', r"Góc hợp bởi tia sáng xa nhất $AD$ và bức tường thẳng đứng $AB$ là:"),
        ('display_math', r"\widehat{BAD} = \widehat{BAC} + \widehat{CAD} \approx 30.964^\circ + 30^\circ = 60.964^\circ."),
        ('bullet', r"Trong tam giác vuông $ABD$ vuông tại $B$:"),
        ('display_math', r"BD = AB \cdot \tan(\widehat{BAD}) \approx 5 \cdot \tan(60.964^\circ) \approx 5 \cdot 1.80125 \approx 9.006\text{ m}."),
        ('p', r"**3. Chiều dài vùng được chiếu sáng:**"),
        ('display_math', r"CD = BD - BC \approx 9.006 - 3 = 6.006\text{ m} \approx 6.0\text{ m}."),
        ('bullet', r"**Kết luận:** Độ dài vùng được chiếu sáng trên mặt đất là xấp xỉ **6.0 mét**.")
    ])

    # =========================================================================
    # PHẦN V. TỰ ĐÁNH GIÁ & NHẬN XÉT CỦA GIÁO VIÊN
    # =========================================================================
    builder.add_self_evaluation([
        "Nhận biết & giải bất phương trình, hệ bất phương trình bậc nhất hai ẩn",
        "Biểu diễn miền nghiệm trên mặt phẳng tọa độ Oxy và giải bài toán quy hoạch tối ưu thực tế",
        "Giá trị lượng giác từ 0° đến 180°, các hệ thức cơ bản, góc bù và dấu lượng giác",
        "Vận dụng định lý sin, định lý cosin, các công thức tính diện tích tam giác",
        "Mô hình hóa hình học và giải bài toán thực tế (đo chiều cao cây, dải sáng chiếu tường)"
    ])

    builder.add_teacher_feedback_and_signature("Thầy Nguyễn Văn Dũng – MathPlus Academy")

    builder.save(output_path)
    print(f"File MathPlus hoàn tất lưu tại: {output_path}")

if __name__ == "__main__":
    out_docx = r"C:\code\mathplus\lop10\MathPlus_Giai_Chi_Tiet_De_Giua_HK1_Toan_10_THPT_Tran_Nhan_Tong_MaDe_294.docx"
    build_mathplus_solution_doc(out_docx)
