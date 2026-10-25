#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script chuyển đổi 3 ảnh đề thi giữa kì 1 Toán 10 THPT Trần Nhân Tông (Mã đề 294)
sang tài liệu Microsoft Word (.docx) chuẩn chỉnh:
- 100% công thức toán OMML (Word native)
- Không có chữ viết tay của học sinh
- Đầy đủ 4 hình vẽ minh họa toán học sắc nét
- Định dạng bảng căn chỉnh các phương án A, B, C, D ngay ngắn
- Ô điền đáp án dạng lưới (grid-in) cho Phần III
"""

import os
import sys
import re

# Thêm đường dẫn tới omml_converter
sys.path.append(r"C:\Users\nguye\.gemini\config\skills\mathplus-giai-toan\scripts")
from omml_converter import latex_to_omml

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

COLOR_BLACK = RGBColor(0, 0, 0)
FONT_NAME = "Times New Roman"

def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_formatted_text(p, text, font_name=FONT_NAME, font_size=Pt(11), 
                       font_color=COLOR_BLACK, bold=False, italic=False):
    """
    Phân tích chuỗi text chứa markdown bold (**...**), italic (*...*) và công thức toán ($...$)
    chuyển đổi thành OMML và text run tương ứng.
    """
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\$[^$]+\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            parts = re.split(r'(\$[^$]+\$)', inner)
            for part in parts:
                if part.startswith('$') and part.endswith('$'):
                    math_code = part[1:-1].strip()
                    try:
                        omml = latex_to_omml(math_code, display=False)
                        p._p.append(omml)
                    except Exception:
                        r = p.add_run(part)
                        r.font.name = font_name
                        r.font.size = font_size
                        r.font.bold = True
                        r.font.color.rgb = font_color
                elif part:
                    r = p.add_run(part)
                    r.font.name = font_name
                    r.font.size = font_size
                    r.font.bold = True
                    r.font.color.rgb = font_color
        elif token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            parts = re.split(r'(\$[^$]+\$)', inner)
            for part in parts:
                if part.startswith('$') and part.endswith('$'):
                    math_code = part[1:-1].strip()
                    try:
                        omml = latex_to_omml(math_code, display=False)
                        p._p.append(omml)
                    except Exception:
                        r = p.add_run(part)
                        r.font.name = font_name
                        r.font.size = font_size
                        r.font.italic = True
                        r.font.color.rgb = font_color
                elif part:
                    r = p.add_run(part)
                    r.font.name = font_name
                    r.font.size = font_size
                    r.font.italic = True
                    r.font.color.rgb = font_color
        elif token.startswith('$') and token.endswith('$'):
            math_code = token[1:-1].strip()
            try:
                omml = latex_to_omml(math_code, display=False)
                p._p.append(omml)
            except Exception:
                r = p.add_run(token)
                r.font.name = font_name
                r.font.size = font_size
                r.font.color.rgb = font_color
                r.font.bold = bold
                r.font.italic = italic
        else:
            r = p.add_run(token)
            r.font.name = font_name
            r.font.size = font_size
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = font_color

def build_exam_document(output_path):
    doc = Document()

    # Thiết lập lề trang
    for section in doc.sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Trường THPT Trần Nhân Tông – Đề kiểm tra giữa học kì 1 Toán 10 (Mã đề 294)")
        hrun.font.name = FONT_NAME
        hrun.font.size = Pt(9)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(100, 100, 100)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Trang ")
        frun.font.name = FONT_NAME
        frun.font.size = Pt(10)
        # Thêm field PAGE
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fldSimple)
        frun2 = fp.add_run(" / 3 - Mã đề 294")
        frun2.font.name = FONT_NAME
        frun2.font.size = Pt(10)

    # 1. HEADER BÀI THI (2 CỘT)
    table_head = doc.add_table(rows=1, cols=2)
    table_head.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_head.autofit = False
    
    # Cột trái
    cell_l = table_head.cell(0, 0)
    cell_l.width = Mm(85)
    p_l1 = cell_l.paragraphs[0]
    p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l1.paragraph_format.space_before = Pt(0)
    p_l1.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_l1, "**SỞ GIÁO DỤC VÀ ĐÀO TẠO HÀ NỘI**\n**TRƯỜNG THPT TRẦN NHÂN TÔNG**\n*(Đề thi có 3 trang)*")
    
    # Cột phải
    cell_r = table_head.cell(0, 1)
    cell_r.width = Mm(85)
    p_r1 = cell_r.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.space_before = Pt(0)
    p_r1.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_r1, "**KIỂM TRA GIỮA HỌC KÌ 1**\n**NĂM HỌC 2026 – 2027**\n**Môn: Toán 10**\n*Thời gian: 90 phút (không kể phát đề)*\n**Mã đề 294**")
    
    # Dòng thông tin học sinh
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(6)
    p_info.paragraph_format.space_after = Pt(8)
    add_formatted_text(p_info, "Họ và tên học sinh: ............................................................................ Số báo danh: ...........................................")

    # =========================================================================
    # PHẦN I. TRẮC NGHIỆM
    # =========================================================================
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(8)
    p_p1.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p1, "**PHẦN I. Trắc nghiệm.** Học sinh trả lời từ câu 1 đến câu 12.")
    
    p_p1_sub = doc.add_paragraph()
    p_p1_sub.paragraph_format.space_before = Pt(0)
    p_p1_sub.paragraph_format.space_after = Pt(6)
    add_formatted_text(p_p1_sub, "*Mỗi câu hỏi học sinh chỉ chọn một phương án.*")

    # --- Câu 1 ---
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(4)
    p_c1.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c1, "**Câu 1.** Bất phương trình nào sau đây **không** là bất phương trình bậc nhất hai ẩn?")
    
    t1 = doc.add_table(rows=2, cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False
    for r in range(2):
        for c in range(2):
            t1.cell(r, c).width = Mm(85)
            set_cell_margins(t1.cell(r, c), 20, 20, 50, 50)
    add_formatted_text(t1.cell(0, 0).paragraphs[0], "**A.** $x - 4y + 2 > 0.$")
    add_formatted_text(t1.cell(0, 1).paragraphs[0], "**B.** $3x + 2y \\le 5.$")
    add_formatted_text(t1.cell(1, 0).paragraphs[0], "**C.** $xy - 2x + 1 < 0.$")
    add_formatted_text(t1.cell(1, 1).paragraphs[0], "**D.** $y + 7 \\ge 0.$")

    # --- Câu 2 ---
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(6)
    p_c2.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c2, "**Câu 2.** Bất phương trình $5x - 2(y - 2x - 3) > 0$ tương đương với")
    
    t2 = doc.add_table(rows=2, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    for r in range(2):
        for c in range(2):
            t2.cell(r, c).width = Mm(85)
            set_cell_margins(t2.cell(r, c), 20, 20, 50, 50)
    add_formatted_text(t2.cell(0, 0).paragraphs[0], "**A.** $9x - 2y + 6 > 0.$")
    add_formatted_text(t2.cell(0, 1).paragraphs[0], "**B.** $x - 2y + 6 > 0.$")
    add_formatted_text(t2.cell(1, 0).paragraphs[0], "**C.** $9x - 2y - 6 > 0.$")
    add_formatted_text(t2.cell(1, 1).paragraphs[0], "**D.** $7x - 2y + 3 > 0.$")

    # --- Câu 3 ---
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.space_before = Pt(6)
    p_c3.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c3, "**Câu 3.** Cặp số nào sau đây **không** là nghiệm của bất phương trình $2x + 3y - 6 > 0$?")
    
    t3 = doc.add_table(rows=1, cols=4)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3.autofit = False
    for c in range(4):
        t3.cell(0, c).width = Mm(42.5)
        set_cell_margins(t3.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t3.cell(0, 0).paragraphs[0], "**A.** $(2; 1).$")
    add_formatted_text(t3.cell(0, 1).paragraphs[0], "**B.** $(0; 3).$")
    add_formatted_text(t3.cell(0, 2).paragraphs[0], "**C.** $(3; 1).$")
    add_formatted_text(t3.cell(0, 3).paragraphs[0], "**D.** $(0; 0).$")

    # --- Câu 4 (Có hình vẽ) ---
    p_c4 = doc.add_paragraph()
    p_c4.paragraph_format.space_before = Pt(6)
    p_c4.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c4, "**Câu 4.** Miền không bị gạch chéo (không kể bờ $d$) trong hình là miền nghiệm của bất phương trình nào?")
    
    # Bảng 2 cột: Cột trái để trống hoặc phương án, cột phải hình vẽ
    t4_wrap = doc.add_table(rows=1, cols=2)
    t4_wrap.alignment = WD_TABLE_ALIGNMENT.CENTER
    t4_wrap.autofit = False
    t4_wrap.cell(0, 0).width = Mm(70)
    t4_wrap.cell(0, 1).width = Mm(100)
    set_cell_margins(t4_wrap.cell(0, 0), 20, 20, 30, 30)
    set_cell_margins(t4_wrap.cell(0, 1), 20, 20, 30, 30)
    
    # Phương án cột trái
    p_opt = t4_wrap.cell(0, 0).paragraphs[0]
    p_opt.paragraph_format.space_before = Pt(4)
    p_opt.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_opt, "**A.** $x + 4y - 8 < 0.$\n\n**B.** $x + 4y - 8 > 0.$\n\n**C.** $4x + y - 8 < 0.$\n\n**D.** $4x + y - 8 > 0.$")
    
    # Hình ảnh cột phải
    p_fig = t4_wrap.cell(0, 1).paragraphs[0]
    p_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig4_path = r"C:\code\mathplus\lop10\figures_de294\c4_oxy.png"
    if os.path.exists(fig4_path):
        p_fig.add_run().add_picture(fig4_path, width=Inches(3.3))

    # --- Câu 5 ---
    p_c5 = doc.add_paragraph()
    p_c5.paragraph_format.space_before = Pt(6)
    p_c5.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c5, "**Câu 5.** Hệ nào sau đây **không** là hệ bất phương trình bậc nhất hai ẩn?")
    
    t5 = doc.add_table(rows=2, cols=2)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    t5.autofit = False
    for r in range(2):
        for c in range(2):
            t5.cell(r, c).width = Mm(85)
            set_cell_margins(t5.cell(r, c), 20, 20, 50, 50)
    add_formatted_text(t5.cell(0, 0).paragraphs[0], "**A.** $\\begin{cases} x + y \\le 5 \\\\ 2x - y > 0 \\end{cases}$")
    add_formatted_text(t5.cell(0, 1).paragraphs[0], "**B.** $\\begin{cases} x^3 - y < 2 \\\\ x + 4y \\le 9 \\end{cases}$")
    add_formatted_text(t5.cell(1, 0).paragraphs[0], "**C.** $\\begin{cases} x \\ge -1 \\\\ 3x + 2y < 7 \\end{cases}$")
    add_formatted_text(t5.cell(1, 1).paragraphs[0], "**D.** $\\begin{cases} 4x - y \\ge 3 \\\\ y \\le 5 \\end{cases}$")

    # --- Câu 6 ---
    p_c6 = doc.add_paragraph()
    p_c6.paragraph_format.space_before = Pt(6)
    p_c6.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c6, "**Câu 6.** Cặp số nào sau đây **không** là nghiệm của hệ $\\begin{cases} x + y - 3 \\le 0 \\\\ 2x - y + 1 > 0 \\end{cases}$ ?")
    
    t6 = doc.add_table(rows=1, cols=4)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    t6.autofit = False
    for c in range(4):
        t6.cell(0, c).width = Mm(42.5)
        set_cell_margins(t6.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t6.cell(0, 0).paragraphs[0], "**A.** $(0; 0).$")
    add_formatted_text(t6.cell(0, 1).paragraphs[0], "**B.** $(1; 1).$")
    add_formatted_text(t6.cell(0, 2).paragraphs[0], "**C.** $(-1; 1).$")
    add_formatted_text(t6.cell(0, 3).paragraphs[0], "**D.** $(2; 0).$")

    # --- Câu 7 ---
    p_c7 = doc.add_paragraph()
    p_c7.paragraph_format.space_before = Pt(6)
    p_c7.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c7, "**Câu 7.** Cho hệ $\\begin{cases} x - \\dfrac{1}{4}y < 3 \\quad (1) \\\\ 4x - y < 8 \\quad (2) \\end{cases}$. Khẳng định đúng về tập nghiệm $S_1, S_2$ của $(1),(2)$ là")
    
    t7 = doc.add_table(rows=1, cols=4)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    t7.autofit = False
    for c in range(4):
        t7.cell(0, c).width = Mm(42.5)
        set_cell_margins(t7.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t7.cell(0, 0).paragraphs[0], "**A.** $S_1 \\subset S_2.$")
    add_formatted_text(t7.cell(0, 1).paragraphs[0], "**B.** $S_2 \\subset S_1.$")
    add_formatted_text(t7.cell(0, 2).paragraphs[0], "**C.** $S_1 = S_2.$")
    add_formatted_text(t7.cell(0, 3).paragraphs[0], "**D.** $S_2 \\ne S_1.$")

    # --- Câu 8 ---
    p_c8 = doc.add_paragraph()
    p_c8.paragraph_format.space_before = Pt(6)
    p_c8.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c8, "**Câu 8.** Cho $\\alpha$ là góc tù và $\\sin\\alpha = \\dfrac{9}{41}$. Tính $\\cos\\alpha$.")
    
    t8 = doc.add_table(rows=1, cols=4)
    t8.alignment = WD_TABLE_ALIGNMENT.CENTER
    t8.autofit = False
    for c in range(4):
        t8.cell(0, c).width = Mm(42.5)
        set_cell_margins(t8.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t8.cell(0, 0).paragraphs[0], "**A.** $-\\dfrac{40}{41}.$")
    add_formatted_text(t8.cell(0, 1).paragraphs[0], "**B.** $\\dfrac{40}{41}.$")
    add_formatted_text(t8.cell(0, 2).paragraphs[0], "**C.** $-\\dfrac{32}{41}.$")
    add_formatted_text(t8.cell(0, 3).paragraphs[0], "**D.** $\\dfrac{32}{41}.$")

    # --- Câu 9 ---
    p_c9 = doc.add_paragraph()
    p_c9.paragraph_format.space_before = Pt(6)
    p_c9.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c9, "**Câu 9.** Cho $\\sin\\alpha = \\dfrac{20}{29}, 90^\\circ < \\alpha < 180^\\circ$. Tính $\\tan(180^\\circ - \\alpha)$.")
    
    t9 = doc.add_table(rows=1, cols=4)
    t9.alignment = WD_TABLE_ALIGNMENT.CENTER
    t9.autofit = False
    for c in range(4):
        t9.cell(0, c).width = Mm(42.5)
        set_cell_margins(t9.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t9.cell(0, 0).paragraphs[0], "**A.** $\\dfrac{20}{21}.$")
    add_formatted_text(t9.cell(0, 1).paragraphs[0], "**B.** $-\\dfrac{20}{21}.$")
    add_formatted_text(t9.cell(0, 2).paragraphs[0], "**C.** $\\dfrac{21}{20}.$")
    add_formatted_text(t9.cell(0, 3).paragraphs[0], "**D.** $-\\dfrac{21}{20}.$")

    # --- Câu 10 ---
    p_c10 = doc.add_paragraph()
    p_c10.paragraph_format.space_before = Pt(6)
    p_c10.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c10, "**Câu 10.** Cho tam giác $ABC$ có $b = 8, c = 6, \\cos A = \\dfrac{3}{4}$. Tính $a$.")
    
    t10 = doc.add_table(rows=1, cols=4)
    t10.alignment = WD_TABLE_ALIGNMENT.CENTER
    t10.autofit = False
    for c in range(4):
        t10.cell(0, c).width = Mm(42.5)
        set_cell_margins(t10.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t10.cell(0, 0).paragraphs[0], "**A.** $2\\sqrt{7}.$")
    add_formatted_text(t10.cell(0, 1).paragraphs[0], "**B.** $4.$")
    add_formatted_text(t10.cell(0, 2).paragraphs[0], "**C.** $\\sqrt{22}.$")
    add_formatted_text(t10.cell(0, 3).paragraphs[0], "**D.** $6.$")

    # --- Câu 11 ---
    p_c11 = doc.add_paragraph()
    p_c11.paragraph_format.space_before = Pt(6)
    p_c11.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c11, "**Câu 11.** Cho tam giác $ABC$ và bán kính ngoại tiếp $R$. Khẳng định nào sau đây sai?")
    
    t11 = doc.add_table(rows=2, cols=2)
    t11.alignment = WD_TABLE_ALIGNMENT.CENTER
    t11.autofit = False
    for r in range(2):
        for c in range(2):
            t11.cell(r, c).width = Mm(85)
            set_cell_margins(t11.cell(r, c), 20, 20, 50, 50)
    add_formatted_text(t11.cell(0, 0).paragraphs[0], "**A.** $\\sin A = \\dfrac{a}{2R}.$")
    add_formatted_text(t11.cell(0, 1).paragraphs[0], "**B.** $\\dfrac{c}{\\sin C} = 2R.$")
    add_formatted_text(t11.cell(1, 0).paragraphs[0], "**C.** $a\\sin B = b\\sin A.$")
    add_formatted_text(t11.cell(1, 1).paragraphs[0], "**D.** $c\\sin C = 2R.$")

    # --- Câu 12 ---
    p_c12 = doc.add_paragraph()
    p_c12.paragraph_format.space_before = Pt(6)
    p_c12.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_c12, "**Câu 12.** Cho tam giác $ABC$ có $a = 10, c = 6, B = 30^\\circ$. Diện tích tam giác bằng")
    
    t12 = doc.add_table(rows=1, cols=4)
    t12.alignment = WD_TABLE_ALIGNMENT.CENTER
    t12.autofit = False
    for c in range(4):
        t12.cell(0, c).width = Mm(42.5)
        set_cell_margins(t12.cell(0, c), 20, 20, 40, 40)
    add_formatted_text(t12.cell(0, 0).paragraphs[0], "**A.** $15.$")
    add_formatted_text(t12.cell(0, 1).paragraphs[0], "**B.** $30.$")
    add_formatted_text(t12.cell(0, 2).paragraphs[0], "**C.** $15\\sqrt{3}.$")
    add_formatted_text(t12.cell(0, 3).paragraphs[0], "**D.** $10\\sqrt{3}.$")

    # =========================================================================
    # PHẦN II. TRẮC NGHIỆM ĐÚNG SAI
    # =========================================================================
    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(12)
    p_p2.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p2, "**PHẦN II. Trắc nghiệm đúng sai.** Học sinh trả lời từ câu 1 đến câu 2.")
    
    p_p2_sub = doc.add_paragraph()
    p_p2_sub.paragraph_format.space_before = Pt(0)
    p_p2_sub.paragraph_format.space_after = Pt(6)
    add_formatted_text(p_p2_sub, "*Trong mỗi ý a), b), c) và d) ở mỗi câu, học sinh chọn đúng hoặc sai.*")

    # --- Câu 1 Phần II ---
    p_p2_c1 = doc.add_paragraph()
    p_p2_c1.paragraph_format.space_before = Pt(4)
    p_p2_c1.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p2_c1, "**Câu 1.** Cho tam giác $ABC$ có $\\widehat{A} = 60^\\circ, AB = 9, AC = 15$. Xét các mệnh đề:")
    
    t_p2_1 = doc.add_table(rows=2, cols=2)
    t_p2_1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_p2_1.autofit = False
    for r in range(2):
        for c in range(2):
            t_p2_1.cell(r, c).width = Mm(85)
            set_cell_margins(t_p2_1.cell(r, c), 20, 20, 50, 50)
    add_formatted_text(t_p2_1.cell(0, 0).paragraphs[0], "**a)** $S = \\dfrac{1}{2}AB \\cdot AC \\cos A.$")
    add_formatted_text(t_p2_1.cell(0, 1).paragraphs[0], "**b)** $S = \\dfrac{135\\sqrt{3}}{4}.$")
    add_formatted_text(t_p2_1.cell(1, 0).paragraphs[0], "**c)** $h_b = 9\\sqrt{3}.$")
    add_formatted_text(t_p2_1.cell(1, 1).paragraphs[0], "**d)** $R = \\sqrt{57}.$")

    # --- Câu 2 Phần II ---
    p_p2_c2 = doc.add_paragraph()
    p_p2_c2.paragraph_format.space_before = Pt(6)
    p_p2_c2.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p2_c2, "**Câu 2.** Cho tam giác $ABC$ có ba cạnh $5, 12, 13$. Xét các mệnh đề:")
    
    t_p2_2 = doc.add_table(rows=2, cols=2)
    t_p2_2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_p2_2.autofit = False
    for r in range(2):
        for c in range(2):
            t_p2_2.cell(r, c).width = Mm(85)
            set_cell_margins(t_p2_2.cell(r, c), 20, 20, 50, 50)
    add_formatted_text(t_p2_2.cell(0, 0).paragraphs[0], "**a)** $S = 30.$")
    add_formatted_text(t_p2_2.cell(0, 1).paragraphs[0], "**b)** $p = 15.$")
    add_formatted_text(t_p2_2.cell(1, 0).paragraphs[0], "**c)** $r = 3.$")
    add_formatted_text(t_p2_2.cell(1, 1).paragraphs[0], "**d)** $R = \\dfrac{13}{2}.$")

    # =========================================================================
    # PHẦN III. TRẢ LỜI NGẮN
    # =========================================================================
    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(12)
    p_p3.paragraph_format.space_after = Pt(6)
    add_formatted_text(p_p3, "**PHẦN III. Trả lời ngắn.** Học sinh trả lời từ câu 1 đến câu 4.")

    def add_answer_grid(container_p_or_cell):
        """Thêm 4 ô vuông kẻ viền đen để điền đáp án dạng grid-in."""
        tbl = doc.add_table(rows=1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.RIGHT
        tbl.autofit = False
        for c in range(4):
            cell = tbl.cell(0, c)
            cell.width = Mm(7)
            # Viền đen 4 phía
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                f'<w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                f'<w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(" ")
            r.font.size = Pt(11)
        return tbl

    # --- Câu 1 Phần III (Có hình vẽ) ---
    t_p3_1_wrap = doc.add_table(rows=1, cols=2)
    t_p3_1_wrap.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_p3_1_wrap.autofit = False
    t_p3_1_wrap.cell(0, 0).width = Mm(100)
    t_p3_1_wrap.cell(0, 1).width = Mm(70)
    set_cell_margins(t_p3_1_wrap.cell(0, 0), 20, 20, 20, 20)
    set_cell_margins(t_p3_1_wrap.cell(0, 1), 20, 20, 20, 20)

    p_p3_c1_txt = t_p3_1_wrap.cell(0, 0).paragraphs[0]
    p_p3_c1_txt.paragraph_format.space_before = Pt(4)
    p_p3_c1_txt.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_p3_c1_txt, "**Câu 1.** Từ vị trí $A$, người ta quan sát thấy một cây cao, biết $AH = 4\\text{ m}, HB = 15\\text{ m}, \\widehat{BAC} = 50^\\circ$. Khi đó chiều cao của cây là bao nhiêu? *(Lấy kết quả làm tròn một chữ số thập phân sau dấu phẩy)*.")
    
    # Ô trả lời trong cell 0
    p_box1 = t_p3_1_wrap.cell(0, 0).add_paragraph()
    p_box1.paragraph_format.space_before = Pt(4)
    p_box1.paragraph_format.space_after = Pt(4)
    p_box1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_formatted_text(p_box1, "*Đáp số:* ")
    # Thêm bảng 4 ô vào cột trái
    grid1 = t_p3_1_wrap.cell(0, 0).add_table(rows=1, cols=4)
    grid1.alignment = WD_TABLE_ALIGNMENT.LEFT
    for c in range(4):
        cell = grid1.cell(0, c)
        cell.width = Mm(7)
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        r = cp.add_run(" ")
        r.font.size = Pt(11)

    # Hình cây cột phải
    p_p3_c1_fig = t_p3_1_wrap.cell(0, 1).paragraphs[0]
    p_p3_c1_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_tree_path = r"C:\code\mathplus\lop10\figures_de294\p3_c1_tree.png"
    if os.path.exists(fig_tree_path):
        p_p3_c1_fig.add_run().add_picture(fig_tree_path, width=Inches(2.7))

    # --- Câu 2 Phần III ---
    p_p3_c2 = doc.add_paragraph()
    p_p3_c2.paragraph_format.space_before = Pt(6)
    p_p3_c2.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p3_c2, "**Câu 2.** Một cửa hàng dự định kinh doanh bổ sung hai dòng sản phẩm dòng M và dòng N, với số vốn ban đầu không vượt quá 39 tỉ đồng. Giá nhập một sản phẩm dòng M là 50 triệu đồng, lợi nhuận dự kiến là 6 triệu đồng; giá nhập một sản phẩm dòng N là 30 triệu đồng, lợi nhuận dự kiến là 3 triệu đồng. Tổng nhu cầu thị trường không vượt quá 1000 sản phẩm và nhu cầu dòng M không vượt quá 2 lần nhu cầu dòng N. Hỏi lợi nhuận lớn nhất có thể thu được là bao nhiêu *(đơn vị: triệu đồng)*?")
    
    # Ô điền đáp án
    p_box2_lead = doc.add_paragraph()
    p_box2_lead.paragraph_format.space_before = Pt(2)
    p_box2_lead.paragraph_format.space_after = Pt(2)
    p_box2_lead.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_formatted_text(p_box2_lead, "*Đáp số:* ")
    add_answer_grid(doc)

    # --- Câu 3 Phần III ---
    p_p3_c3 = doc.add_paragraph()
    p_p3_c3.paragraph_format.space_before = Pt(6)
    p_p3_c3.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p3_c3, "**Câu 3.** Một hộ nông dân có $1\\,200\\text{ m}^2$ đất. Trên mỗi $100\\text{ m}^2$, loại cây thứ nhất cần 20 công và thu 5 triệu đồng; loại thứ hai cần 30 công và thu 7 triệu đồng. Tổng số công không quá 300. Hỏi thu nhập lớn nhất là bao nhiêu triệu đồng?")
    
    p_box3_lead = doc.add_paragraph()
    p_box3_lead.paragraph_format.space_before = Pt(2)
    p_box3_lead.paragraph_format.space_after = Pt(2)
    p_box3_lead.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_formatted_text(p_box3_lead, "*Đáp số:* ")
    add_answer_grid(doc)

    # --- Câu 4 Phần III ---
    p_p3_c4 = doc.add_paragraph()
    p_p3_c4.paragraph_format.space_before = Pt(6)
    p_p3_c4.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_p3_c4, "**Câu 4.** Cho tam giác $ABC$ có $BC = 9, \\widehat{B} = 60^\\circ, \\widehat{C} = 75^\\circ$. Tính $AC$ *(làm tròn đến hàng phần trăm)*.")
    
    p_box4_lead = doc.add_paragraph()
    p_box4_lead.paragraph_format.space_before = Pt(2)
    p_box4_lead.paragraph_format.space_after = Pt(2)
    p_box4_lead.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_formatted_text(p_box4_lead, "*Đáp số:* ")
    add_answer_grid(doc)

    # =========================================================================
    # PHẦN IV. TỰ LUẬN
    # =========================================================================
    p_p4 = doc.add_paragraph()
    p_p4.paragraph_format.space_before = Pt(12)
    p_p4.paragraph_format.space_after = Pt(6)
    add_formatted_text(p_p4, "**PHẦN IV. Tự luận.** Học sinh làm ra giấy từ câu 1 đến câu 3.")

    # --- Câu 1 Phần IV ---
    p_p4_c1 = doc.add_paragraph()
    p_p4_c1.paragraph_format.space_before = Pt(4)
    p_p4_c1.paragraph_format.space_after = Pt(6)
    add_formatted_text(p_p4_c1, "**Câu 1.** Một cơ sở sử dụng nguyên liệu $K, P$. Mỗi tấn $K$ giá 4 triệu đồng, tạo được 30 kg sản phẩm I và 30 kg sản phẩm II; mỗi tấn $P$ giá 5 triệu đồng, tạo được 20 kg sản phẩm I và 40 kg sản phẩm II. Cần ít nhất 90 kg sản phẩm I và 120 kg sản phẩm II, đồng thời mua không quá 4 tấn mỗi loại. Tìm phương án có chi phí nhỏ nhất.")

    # --- Câu 2 Phần IV (Có hình vẽ) ---
    t_p4_2_wrap = doc.add_table(rows=1, cols=2)
    t_p4_2_wrap.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_p4_2_wrap.autofit = False
    t_p4_2_wrap.cell(0, 0).width = Mm(110)
    t_p4_2_wrap.cell(0, 1).width = Mm(60)
    set_cell_margins(t_p4_2_wrap.cell(0, 0), 20, 20, 20, 20)
    set_cell_margins(t_p4_2_wrap.cell(0, 1), 20, 20, 20, 20)

    p_p4_c2_txt = t_p4_2_wrap.cell(0, 0).paragraphs[0]
    p_p4_c2_txt.paragraph_format.space_before = Pt(4)
    p_p4_c2_txt.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_p4_c2_txt, "**Câu 2.** Từ một tấm bìa hình tròn, cắt ra tam giác $ABC$ có $AB = 5\\text{ cm}, AC = 7\\text{ cm}$ và $\\widehat{B} = 60^\\circ$. Tính $BC$ và bán kính $R$ của hình tròn *(làm tròn $R$ đến hàng phần mười)*.")

    p_p4_c2_fig = t_p4_2_wrap.cell(0, 1).paragraphs[0]
    p_p4_c2_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_circle_path = r"C:\code\mathplus\lop10\figures_de294\p4_c2_circle.png"
    if os.path.exists(fig_circle_path):
        p_p4_c2_fig.add_run().add_picture(fig_circle_path, width=Inches(2.2))

    # --- Câu 3 Phần IV (Có hình vẽ) ---
    t_p4_3_wrap = doc.add_table(rows=1, cols=2)
    t_p4_3_wrap.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_p4_3_wrap.autofit = False
    t_p4_3_wrap.cell(0, 0).width = Mm(105)
    t_p4_3_wrap.cell(0, 1).width = Mm(65)
    set_cell_margins(t_p4_3_wrap.cell(0, 0), 20, 20, 20, 20)
    set_cell_margins(t_p4_3_wrap.cell(0, 1), 20, 20, 20, 20)

    p_p4_c3_txt = t_p4_3_wrap.cell(0, 0).paragraphs[0]
    p_p4_c3_txt.paragraph_format.space_before = Pt(4)
    p_p4_c3_txt.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_p4_c3_txt, "**Câu 3.** Người ta lắp một thiết bị chiếu sáng sát tường, cao 5 m so với mặt đất. Góc chiếu sáng là $30^\\circ$ và dải sáng bắt đầu tại điểm cách chân tường 3 m. Độ dài vùng được chiếu sáng trên mặt đất bằng bao nhiêu mét? *(Làm tròn đến hàng phần mười)*.")

    p_p4_c3_fig = t_p4_3_wrap.cell(0, 1).paragraphs[0]
    p_p4_c3_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_light_path = r"C:\code\mathplus\lop10\figures_de294\p4_c3_light.png"
    if os.path.exists(fig_light_path):
        p_p4_c3_fig.add_run().add_picture(fig_light_path, width=Inches(2.4))

    # Kết thúc đề thi
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(14)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_text(p_end, "---------- **HẾT** ----------\n*(Cán bộ coi thi không giải thích gì thêm)*", italic=True)

    # Lưu file
    doc.save(output_path)
    print(f"File Word created successfully at: {output_path}")

if __name__ == "__main__":
    out_docx = r"C:\code\mathplus\lop10\De_Kiem_Tra_Giua_HK1_Toan_10_THPT_Tran_Nhan_Tong_MaDe_294.docx"
    build_exam_document(out_docx)
