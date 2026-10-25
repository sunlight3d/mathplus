#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script tạo tài liệu Word ĐÁP ÁN & LỜI GIẢI CHI TIẾT
Đề kiểm tra giữa học kì 1 Toán 10 THPT Trần Nhân Tông (Mã đề 294)
- 100% công thức toán OMML
- Bảng tổng hợp đáp án nhanh
- Lời giải chi tiết từng bước, lập luận sư phạm chặt chẽ
- Đầy đủ hình vẽ minh họa
"""

import os
import sys
import re

sys.path.append(r"C:\Users\nguye\.gemini\config\skills\mathplus-giai-toan\scripts")
from omml_converter import latex_to_omml

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

COLOR_BLACK = RGBColor(0, 0, 0)
FONT_NAME = "Times New Roman"

def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", sz="4", color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_formatted_text(p, text, font_name=FONT_NAME, font_size=Pt(11), 
                       font_color=COLOR_BLACK, bold=False, italic=False):
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\$[^$]+\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            parts = re.split(r'(\$[^$]+\$)', inner)
            for part in parts:
                if part.startswith('$') and part.endswith('$'):
                    try:
                        p._p.append(latex_to_omml(part[1:-1].strip(), display=False))
                    except Exception:
                        r = p.add_run(part)
                        r.font.name = font_name; r.font.size = font_size; r.font.bold = True; r.font.color.rgb = font_color
                elif part:
                    r = p.add_run(part)
                    r.font.name = font_name; r.font.size = font_size; r.font.bold = True; r.font.color.rgb = font_color
        elif token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            parts = re.split(r'(\$[^$]+\$)', inner)
            for part in parts:
                if part.startswith('$') and part.endswith('$'):
                    try:
                        p._p.append(latex_to_omml(part[1:-1].strip(), display=False))
                    except Exception:
                        r = p.add_run(part)
                        r.font.name = font_name; r.font.size = font_size; r.font.italic = True; r.font.color.rgb = font_color
                elif part:
                    r = p.add_run(part)
                    r.font.name = font_name; r.font.size = font_size; r.font.italic = True; r.font.color.rgb = font_color
        elif token.startswith('$') and token.endswith('$'):
            try:
                p._p.append(latex_to_omml(token[1:-1].strip(), display=False))
            except Exception:
                r = p.add_run(token)
                r.font.name = font_name; r.font.size = font_size; r.font.color.rgb = font_color; r.font.bold = bold; r.font.italic = italic
        else:
            r = p.add_run(token)
            r.font.name = font_name; r.font.size = font_size; r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = font_color

def build_solution_doc(output_path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Mm(20); s.bottom_margin = Mm(20); s.left_margin = Mm(20); s.right_margin = Mm(20)
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("THPT Trần Nhân Tông – Đáp án & Lời giải chi tiết Giữa HK1 Toán 10 (Mã đề 294)")
        hrun.font.name = FONT_NAME; hrun.font.size = Pt(9); hrun.font.italic = True; hrun.font.color.rgb = RGBColor(100, 100, 100)

        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Trang ")
        frun.font.name = FONT_NAME; frun.font.size = Pt(10)
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fldSimple)

    # Tiêu đề
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_title, "**ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT**\n**ĐỀ KIỂM TRA GIỮA HỌC KÌ 1 – NĂM HỌC 2026 – 2027**\n**MÔN: TOÁN 10 – MÃ ĐỀ: 294**\n**TRƯỜNG THPT TRẦN NHÂN TÔNG – HÀ NỘI**")

    # =========================================================================
    # BẢNG ĐÁP ÁN NHANH
    # =========================================================================
    p_tab_t = doc.add_paragraph()
    p_tab_t.paragraph_format.space_before = Pt(8); p_tab_t.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_tab_t, "**I. BẢNG ĐÁP ÁN NHANH**")

    # Bảng Phần I (12 câu)
    p_ans1 = doc.add_paragraph()
    p_ans1.paragraph_format.space_before = Pt(2); p_ans1.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_ans1, "**1. PHẦN I - Trắc nghiệm nhiều lựa chọn:**")
    
    t_ans1 = doc.add_table(rows=2, cols=12)
    t_ans1.alignment = WD_TABLE_ALIGNMENT.CENTER
    ans1_keys = ["C", "A", "D", "A", "B", "C", "B", "A", "A", "A", "D", "A"]
    for c in range(12):
        cell_q = t_ans1.cell(0, c)
        cell_a = t_ans1.cell(1, c)
        set_cell_borders(cell_q); set_cell_borders(cell_a)
        
        pq = cell_q.paragraphs[0]; pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pq.paragraph_format.space_before = Pt(2); pq.paragraph_format.space_after = Pt(2)
        add_formatted_text(pq, f"**{c+1}**")
        
        pa = cell_a.paragraphs[0]; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pa.paragraph_format.space_before = Pt(2); pa.paragraph_format.space_after = Pt(2)
        add_formatted_text(pa, f"**{ans1_keys[c]}**")

    # Bảng Phần II (Đúng / Sai)
    p_ans2 = doc.add_paragraph()
    p_ans2.paragraph_format.space_before = Pt(6); p_ans2.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_ans2, "**2. PHẦN II - Trắc nghiệm đúng sai:**")
    
    t_ans2 = doc.add_table(rows=3, cols=5)
    t_ans2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ["Câu", "Lệnh hỏi a", "Lệnh hỏi b", "Lệnh hỏi c", "Lệnh hỏi d"]
    for j, h in enumerate(headers2):
        cell = t_ans2.cell(0, j); set_cell_borders(cell)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_formatted_text(p, f"**{h}**")
        
    p2_data = [
        ("Câu 1", "S", "Đ", "S", "Đ"),
        ("Câu 2", "Đ", "Đ", "S", "Đ")
    ]
    for i, row in enumerate(p2_data):
        for j, val in enumerate(row):
            cell = t_ans2.cell(i+1, j); set_cell_borders(cell)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_text(p, f"**{val}**" if j > 0 else val)

    # Bảng Phần III (Trả lời ngắn)
    p_ans3 = doc.add_paragraph()
    p_ans3.paragraph_format.space_before = Pt(6); p_ans3.paragraph_format.space_after = Pt(2)
    add_formatted_text(p_ans3, "**3. PHẦN III - Trả lời ngắn:**")
    
    t_ans3 = doc.add_table(rows=2, cols=4)
    t_ans3.alignment = WD_TABLE_ALIGNMENT.CENTER
    ans3_keys = [("Câu 1", "14.5"), ("Câu 2", "4500"), ("Câu 3", "72"), ("Câu 4", "11.02")]
    for j, (q, a) in enumerate(ans3_keys):
        cq = t_ans3.cell(0, j); ca = t_ans3.cell(1, j)
        set_cell_borders(cq); set_cell_borders(ca)
        pq = cq.paragraphs[0]; pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_formatted_text(pq, f"**{q}**")
        pa = ca.paragraphs[0]; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_formatted_text(pa, f"**{a}**")

    # =========================================================================
    # LỜI GIẢI CHI TIẾT
    # =========================================================================
    p_detail_title = doc.add_paragraph()
    p_detail_title.paragraph_format.space_before = Pt(14); p_detail_title.paragraph_format.space_after = Pt(6)
    add_formatted_text(p_detail_title, "**II. LỜI GIẢI CHI TIẾT TỪNG CÂU**")

    def add_sol_item(q_num, q_text, sol_text, choice_key=None, fig_path=None, fig_w=Inches(3.0)):
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8); p_q.paragraph_format.space_after = Pt(2)
        prefix = f"**Câu {q_num}.** "
        add_formatted_text(p_q, prefix + q_text)
        
        if fig_path and os.path.exists(fig_path):
            pf = doc.add_paragraph()
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.paragraph_format.space_before = Pt(4); pf.paragraph_format.space_after = Pt(4)
            pf.add_run().add_picture(fig_path, width=fig_w)

        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(2); p_s.paragraph_format.space_after = Pt(6)
        ans_lead = f"**Lời giải:** {sol_text}"
        if choice_key:
            ans_lead += f" $\\Rightarrow$ **Chọn phương án {choice_key}.**"
        add_formatted_text(p_s, ans_lead)

    # PHẦN I
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(10); p_sec1.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_sec1, "**PHẦN I. TRẮC NGHIỆM**")

    add_sol_item(1, "Bất phương trình nào sau đây **không** là bất phương trình bậc nhất hai ẩn?",
                 "Bất phương trình bậc nhất hai ẩn có dạng tổng quát là $ax + by + c < 0$ (hoặc $> 0, \\le 0, \\ge 0$) với $a, b$ không đồng thời bằng 0. "
                 "Ở phương án C, biểu thức $xy - 2x + 1 < 0$ chứa số hạng bậc hai $xy$ nên không là bất phương trình bậc nhất hai ẩn.",
                 "C")

    add_sol_item(2, "Bất phương trình $5x - 2(y - 2x - 3) > 0$ tương đương với",
                 "Khai triển vế trái: $5x - 2y + 4x + 6 > 0 \\iff 9x - 2y + 6 > 0$.",
                 "A")

    add_sol_item(3, "Cặp số nào sau đây **không** là nghiệm của bất phương trình $2x + 3y - 6 > 0$?",
                 "Thay lần lượt các cặp số vào vế trái $f(x, y) = 2x + 3y - 6$:\n"
                 "- Với $(2; 1): 2(2) + 3(1) - 6 = 1 > 0$ (thỏa mãn).\n"
                 "- Với $(0; 3): 2(0) + 3(3) - 6 = 3 > 0$ (thỏa mãn).\n"
                 "- Với $(3; 1): 2(3) + 3(1) - 6 = 3 > 0$ (thỏa mãn).\n"
                 "- Với $(0; 0): 2(0) + 3(0) - 6 = -6 \\not> 0$ (không thỏa mãn).",
                 "D")

    add_sol_item(4, "Miền không bị gạch chéo (không kể bờ $d$) trong hình là miền nghiệm của bất phương trình nào?",
                 "Đường thẳng $d$ cắt trục tung tại $(0; 2)$ và cắt trục hoành tại $(8; 0)$. "
                 "Phương trình đoạn chắn của đường thẳng $d$ là $\\dfrac{x}{8} + \\dfrac{y}{2} = 1 \\iff x + 4y - 8 = 0$. "
                 "Gốc tọa độ $O(0; 0)$ thuộc miền không bị gạch chéo. Thay $x = 0, y = 0$ vào biểu thức $x + 4y - 8$ ta được $0 + 0 - 8 = -8 < 0$. "
                 "Do không kể bờ $d$, bất phương trình cần tìm là $x + 4y - 8 < 0$.",
                 "A",
                 fig_path=r"C:\code\mathplus\lop10\figures_de294\c4_oxy.png", fig_w=Inches(2.8))

    add_sol_item(5, "Hệ nào sau đây **không** là hệ bất phương trình bậc nhất hai ẩn?",
                 "Ở phương án B, hệ $\\begin{cases} x^3 - y < 2 \\\\ x + 4y \\le 9 \\end{cases}$ chứa bất phương trình bậc ba $x^3 - y < 2$ nên không phải là hệ bất phương trình bậc nhất hai ẩn.",
                 "B")

    add_sol_item(6, "Cặp số nào sau đây **không** là nghiệm của hệ $\\begin{cases} x + y - 3 \\le 0 \\\\ 2x - y + 1 > 0 \\end{cases}$?",
                 "Thay cặp số $(-1; 1)$ vào bất phương trình thứ hai của hệ: $2(-1) - 1 + 1 = -2 \\not> 0$ (sai). "
                 "Do đó $(-1; 1)$ không là nghiệm của hệ.",
                 "C")

    add_sol_item(7, "Cho hệ $\\begin{cases} x - \\dfrac{1}{4}y < 3 \\quad (1) \\\\ 4x - y < 8 \\quad (2) \\end{cases}$. Khẳng định đúng về tập nghiệm $S_1, S_2$ của $(1),(2)$ là",
                 "Bất phương trình $(1) \\iff 4x - y < 12$. Miền nghiệm $S_1$ là nửa mặt phẳng bờ $4x - y = 12$ chứa gốc $O(0; 0)$.\n"
                 "Bất phương trình $(2) \\iff 4x - y < 8$. Miền nghiệm $S_2$ là nửa mặt phẳng bờ $4x - y = 8$ chứa gốc $O(0; 0)$.\n"
                 "Vì với mọi điểm $(x, y)$ thỏa mãn $4x - y < 8$ thì hiển nhiên $4x - y < 12$, nên mọi nghiệm của $(2)$ đều là nghiệm của $(1)$. "
                 "Do đó $S_2 \\subset S_1$.",
                 "B")

    add_sol_item(8, "Cho $\\alpha$ là góc tù và $\\sin\\alpha = \\dfrac{9}{41}$. Tính $\\cos\\alpha$.",
                 "Vì $\\alpha$ là góc tù ($90^\\circ < \\alpha < 180^\\circ$) nên $\\cos\\alpha < 0$.\n"
                 "Áp dụng hệ thức cơ bản $\\sin^2\\alpha + \\cos^2\\alpha = 1$, ta có:\n"
                 "$\\cos^2\\alpha = 1 - \\sin^2\\alpha = 1 - \\left(\\dfrac{9}{41}\\right)^2 = \\dfrac{1681 - 81}{1681} = \\dfrac{1600}{1681} = \\left(\\dfrac{40}{41}\\right)^2$.\n"
                 "Do $\\cos\\alpha < 0$ nên $\\cos\\alpha = -\\sqrt{\\dfrac{1600}{1681}} = -\\dfrac{40}{41}$.",
                 "A")

    add_sol_item(9, "Cho $\\sin\\alpha = \\dfrac{20}{29}, 90^\\circ < \\alpha < 180^\\circ$. Tính $\\tan(180^\\circ - \\alpha)$.",
                 "Vì $90^\\circ < \\alpha < 180^\\circ \\implies \\cos\\alpha < 0$.\n"
                 "Ta có $\\cos^2\\alpha = 1 - \\sin^2\\alpha = 1 - \\left(\\dfrac{20}{29}\\right)^2 = \\dfrac{841 - 400}{841} = \\dfrac{441}{841} = \\left(\\dfrac{21}{29}\\right)^2 \\implies \\cos\\alpha = -\\dfrac{21}{29}$.\n"
                 "Suy ra $\\tan\\alpha = \\dfrac{\\sin\\alpha}{\\cos\\alpha} = \\dfrac{20/29}{-21/29} = -\\dfrac{20}{21}$.\n"
                 "Áp dụng công thức hai góc bù nhau: $\\tan(180^\\circ - \\alpha) = -\\tan\\alpha = -\\left(-\\dfrac{20}{21}\\right) = \\dfrac{20}{21}$.",
                 "A")

    add_sol_item(10, "Cho tam giác $ABC$ có $b = 8, c = 6, \\cos A = \\dfrac{3}{4}$. Tính $a$.",
                 "Áp dụng định lí cosin trong tam giác $ABC$:\n"
                 "$a^2 = b^2 + c^2 - 2bc\\cos A = 8^2 + 6^2 - 2 \\cdot 8 \\cdot 6 \\cdot \\dfrac{3}{4} = 64 + 36 - 72 = 28$.\n"
                 "Suy ra $a = \\sqrt{28} = 2\\sqrt{7}$.",
                 "A")

    add_sol_item(11, "Cho tam giác $ABC$ và bán kính ngoại tiếp $R$. Khẳng định nào sau đây sai?",
                 "Theo định lí sin trong tam giác $ABC$: $\\dfrac{a}{\\sin A} = \\dfrac{b}{\\sin B} = \\dfrac{c}{\\sin C} = 2R$.\n"
                 "- Khẳng định A: $\\sin A = \\dfrac{a}{2R}$ (đúng).\n"
                 "- Khẳng định B: $\\dfrac{c}{\\sin C} = 2R$ (đúng).\n"
                 "- Khẳng định C: $\\dfrac{a}{\\sin A} = \\dfrac{b}{\\sin B} \\iff a\\sin B = b\\sin A$ (đúng).\n"
                 "- Khẳng định D: $c\\sin C = 2R$ là khẳng định sai vì $\\dfrac{c}{\\sin C} = 2R$.",
                 "D")

    add_sol_item(12, "Cho tam giác $ABC$ có $a = 10, c = 6, B = 30^\\circ$. Diện tích tam giác bằng",
                 "Áp dụng công thức diện tích tam giác:\n"
                 "$S = \\dfrac{1}{2}ac\\sin B = \\dfrac{1}{2} \\cdot 10 \\cdot 6 \\cdot \\sin 30^\\circ = 30 \\cdot \\dfrac{1}{2} = 15$.",
                 "A")

    # PHẦN II
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(10); p_sec2.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_sec2, "**PHẦN II. TRẮC NGHIỆM ĐÚNG SAI**")

    add_sol_item("1 (Phần II)", "Cho tam giác $ABC$ có $\\widehat{A} = 60^\\circ, AB = 9, AC = 15$. Xét các mệnh đề:",
                 "\n- **Ý a)** $S = \\dfrac{1}{2}AB \\cdot AC \\cos A$: **SAI**. Công thức đúng là $S = \\dfrac{1}{2}AB \\cdot AC \\sin A$.\n"
                 "- **Ý b)** $S = \\dfrac{135\\sqrt{3}}{4}$: **ĐÚNG**. Vì $S = \\dfrac{1}{2} \\cdot 9 \\cdot 15 \\cdot \\sin 60^\\circ = \\dfrac{135}{2} \\cdot \\dfrac{\\sqrt{3}}{2} = \\dfrac{135\\sqrt{3}}{4}$.\n"
                 "- **Ý c)** $h_b = 9\\sqrt{3}$: **SAI**. Ta có $S = \\dfrac{1}{2}b \\cdot h_b = \\dfrac{1}{2}AC \\cdot h_b \\implies h_b = \\dfrac{2S}{AC} = \\dfrac{2 \\cdot \\frac{135\\sqrt{3}}{4}}{15} = \\dfrac{9\\sqrt{3}}{2} \\ne 9\\sqrt{3}$.\n"
                 "- **Ý d)** $R = \\sqrt{57}$: **ĐÚNG**. Áp dụng định lí cosin:\n"
                 "  $a^2 = b^2 + c^2 - 2bc\\cos A = 15^2 + 9^2 - 2 \\cdot 15 \\cdot 9 \\cdot \\cos 60^\\circ = 225 + 81 - 135 = 171 \\implies a = 3\\sqrt{19}$.\n"
                 "  Bán kính $R = \\dfrac{a}{2\\sin A} = \\dfrac{3\\sqrt{19}}{2 \\cdot \\frac{\\sqrt{3}}{2}} = \\dfrac{3\\sqrt{19}}{\\sqrt{3}} = \\sqrt{57}$.")

    add_sol_item("2 (Phần II)", "Cho tam giác $ABC$ có ba cạnh $5, 12, 13$. Xét các mệnh đề:",
                 "Nhận xét: $5^2 + 12^2 = 25 + 144 = 169 = 13^2$, theo định lí Pytago đảo thì tam giác $ABC$ vuông tại đỉnh đối diện cạnh 13.\n"
                 "- **Ý a)** $S = 30$: **ĐÚNG**. Vì tam giác vuông có diện tích bằng nửa tích hai cạnh góc vuông: $S = \\dfrac{1}{2} \\cdot 5 \\cdot 12 = 30$.\n"
                 "- **Ý b)** $p = 15$: **ĐÚNG**. Nửa chu vi $p = \\dfrac{5 + 12 + 13}{2} = 15$.\n"
                 "- **Ý c)** $r = 3$: **SAI**. Bán kính đường tròn nội tiếp $r = \\dfrac{S}{p} = \\dfrac{30}{15} = 2 \\ne 3$.\n"
                 "- **Ý d)** $R = \\dfrac{13}{2}$: **ĐÚNG**. Bán kính đường tròn ngoại tiếp tam giác vuông bằng nửa cạnh huyền: $R = \\dfrac{13}{2} = 6.5$.")

    # PHẦN III
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(10); p_sec3.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_sec3, "**PHẦN III. TRẢ LỜI NGẮN**")

    add_sol_item("1 (Phần III)", "Từ vị trí $A$, người ta quan sát thấy một cây cao, biết $AH = 4\\text{ m}, HB = 15\\text{ m}, \\widehat{BAC} = 50^\\circ$. Khi đó chiều cao của cây là bao nhiêu? *(Lấy kết quả làm tròn một chữ số thập phân sau dấu phẩy)*.",
                 "Kẻ $AK \\perp BC$ tại $K$. Do cây $BC$ vuông góc với mặt đất $HB$ tại $B$ và $AH \\perp HB$, tứ giác $AHBK$ là hình chữ nhật.\n"
                 "Suy ra $AK = HB = 15\\text{ m}$ và $BK = AH = 4\\text{ m}$.\n"
                 "Trong tam giác vuông $AKB$ tại $K$: $\\tan(\\widehat{KAB}) = \\dfrac{BK}{AK} = \\dfrac{4}{15} \\implies \\widehat{KAB} \\approx 14.931^\\circ$.\n"
                 "Tia $AB$ nằm giữa hai tia $AK$ và $AC$, do đó:\n"
                 "$\\widehat{KAC} = \\widehat{BAC} - \\widehat{KAB} = 50^\\circ - 14.931^\\circ = 35.069^\\circ$.\n"
                 "Trong tam giác vuông $AKC$ tại $K$:\n"
                 "$KC = AK \\cdot \\tan(\\widehat{KAC}) = 15 \\cdot \\tan(35.069^\\circ) \\approx 15 \\cdot 0.70196 \\approx 10.529\\text{ m}$.\n"
                 "Chiều cao của cây là: $BC = BK + KC = 4 + 10.529 = 14.529\\text{ m} \\approx 14.5\\text{ m}$.\n"
                 "$\\Rightarrow$ **Đáp số điền ô: 14.5**",
                 fig_path=r"C:\code\mathplus\lop10\figures_de294\p3_c1_tree.png", fig_w=Inches(3.2))

    add_sol_item("2 (Phần III)", "Một cửa hàng dự định kinh doanh bổ sung hai dòng sản phẩm dòng M và dòng N, với số vốn ban đầu không vượt quá 39 tỉ đồng. Giá nhập một sản phẩm dòng M là 50 triệu đồng, lợi nhuận dự kiến là 6 triệu đồng; giá nhập một sản phẩm dòng N là 30 triệu đồng, lợi nhuận dự kiến là 3 triệu đồng. Tổng nhu cầu thị trường không vượt quá 1000 sản phẩm và nhu cầu dòng M không vượt quá 2 lần nhu cầu dòng N. Hỏi lợi nhuận lớn nhất có thể thu được là bao nhiêu *(đơn vị: triệu đồng)*?",
                 "Gọi $x, y$ lần lượt là số sản phẩm dòng M và dòng N cửa hàng nhập về kinh doanh ($x, y \\ge 0, x, y \\in \\mathbb{N}$).\n"
                 "Theo đề bài, ta có hệ bất phương trình ràng buộc:\n"
                 "1) Vốn: $50x + 30y \\le 39\\,000 \\iff 5x + 3y \\le 3\\,900$.\n"
                 "2) Nhu cầu thị trường: $x + y \\le 1\\,000$.\n"
                 "3) Nhu cầu tương quan: $x \\le 2y$.\n"
                 "Lợi nhuận thu được là $F(x, y) = 6x + 3y$ (triệu đồng).\n"
                 "Miền nghiệm của hệ là một đa giác lồi có các đỉnh:\n"
                 "- Gốc tọa độ $O(0; 0) \\implies F(0, 0) = 0$.\n"
                 "- Điểm $A(0; 1000) \\implies F(0, 1000) = 3(1000) = 3\\,000$ triệu đồng.\n"
                 "- Giao điểm $B$ của $x + y = 1000$ và $5x + 3y = 3900$:\n"
                 "  Giải hệ ta được $x = 450, y = 550$ (thỏa mãn $x \\le 2y$). Lợi nhuận $F(450, 550) = 6(450) + 3(550) = 2700 + 1650 = 4\\,350$ triệu đồng.\n"
                 "- Giao điểm $C$ của $5x + 3y = 3900$ và $x = 2y$:\n"
                 "  Thay $x = 2y$ vào: $5(2y) + 3y = 3900 \\iff 13y = 3900 \\implies y = 300, x = 600$ (thỏa mãn $x + y = 900 \\le 1000$).\n"
                 "  Lợi nhuận $F(600, 300) = 6(600) + 3(300) = 3600 + 900 = 4\\,500$ triệu đồng.\n"
                 "So sánh các giá trị tại các đỉnh, lợi nhuận lớn nhất đạt được là **4500 triệu đồng** khi cửa hàng kinh doanh 600 sản phẩm M và 300 sản phẩm N.\n"
                 "$\\Rightarrow$ **Đáp số điền ô: 4500**")

    add_sol_item("3 (Phần III)", "Một hộ nông dân có $1\\,200\\text{ m}^2$ đất. Trên mỗi $100\\text{ m}^2$, loại cây thứ nhất cần 20 công và thu 5 triệu đồng; loại thứ hai cần 30 công và thu 7 triệu đồng. Tổng số công không quá 300. Hỏi thu nhập lớn nhất là bao nhiêu triệu đồng?",
                 "Gọi $x, y$ lần lượt là diện tích đất dùng để trồng loại cây thứ nhất và thứ hai (đơn vị: $100\\text{ m}^2$, điều kiện $x, y \\ge 0$).\n"
                 "Hệ bất phương trình mô tả các điều kiện bài toán:\n"
                 "$\\begin{cases} 100x + 100y \\le 1200 \\\\ 20x + 30y \\le 300 \\\\ x \\ge 0, y \\ge 0 \\end{cases} \\iff \\begin{cases} x + y \\le 12 \\\\ 2x + 3y \\le 30 \\\\ x \\ge 0, y \\ge 0 \\end{cases}$\n"
                 "Tổng thu nhập cần tối đa hóa là $T(x, y) = 5x + 7y$ (triệu đồng).\n"
                 "Các đỉnh của miền nghiệm gồm:\n"
                 "- $O(0; 0) \\implies T = 0$.\n"
                 "- $A(12; 0) \\implies T = 5(12) = 60$ triệu đồng.\n"
                 "- $B(0; 10) \\implies T = 7(10) = 70$ triệu đồng.\n"
                 "- Giao điểm $C$ của $x + y = 12$ và $2x + 3y = 30$:\n"
                 "  Giải hệ: $2(12 - y) + 3y = 30 \\implies y = 6 \\implies x = 6$.\n"
                 "  Thu nhập tại $C(6; 6)$ là $T(6, 6) = 5(6) + 7(6) = 30 + 42 = 72$ triệu đồng.\n"
                 "Vậy thu nhập lớn nhất là **72 triệu đồng**.\n"
                 "$\\Rightarrow$ **Đáp số điền ô: 72**")

    add_sol_item("4 (Phần III)", "Cho tam giác $ABC$ có $BC = 9, \\widehat{B} = 60^\\circ, \\widehat{C} = 75^\\circ$. Tính $AC$ *(làm tròn đến hàng phần trăm)*.",
                 "Tổng ba góc trong tam giác $ABC$ bằng $180^\\circ$, suy ra:\n"
                 "$\\widehat{A} = 180^\\circ - (\\widehat{B} + \\widehat{C}) = 180^\\circ - (60^\\circ + 75^\\circ) = 45^\\circ$.\n"
                 "Áp dụng định lí sin trong tam giác $ABC$:\n"
                 "$\\dfrac{AC}{\\sin B} = \\dfrac{BC}{\\sin A} \\implies AC = BC \\cdot \\dfrac{\\sin B}{\\sin A} = 9 \\cdot \\dfrac{\\sin 60^\\circ}{\\sin 45^\\circ} = 9 \\cdot \\dfrac{\\sqrt{3}/2}{\\sqrt{2}/2} = \\dfrac{9\\sqrt{6}}{2} \\approx 11.0227$.\n"
                 "Làm tròn đến hàng phần trăm ta được $AC \\approx 11.02$.\n"
                 "$\\Rightarrow$ **Đáp số điền ô: 11.02**")

    # PHẦN IV
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(10); p_sec4.paragraph_format.space_after = Pt(4)
    add_formatted_text(p_sec4, "**PHẦN IV. TỰ LUẬN**")

    add_sol_item("1 (Phần IV)", "Một cơ sở sử dụng nguyên liệu $K, P$. Mỗi tấn $K$ giá 4 triệu đồng, tạo được 30 kg sản phẩm I và 30 kg sản phẩm II; mỗi tấn $P$ giá 5 triệu đồng, tạo được 20 kg sản phẩm I và 40 kg sản phẩm II. Cần ít nhất 90 kg sản phẩm I và 120 kg sản phẩm II, đồng thời mua không quá 4 tấn mỗi loại. Tìm phương án có chi phí nhỏ nhất.",
                 "Gọi $x, y$ lần lượt là số tấn nguyên liệu $K$ và $P$ cơ sở cần mua ($0 \\le x \\le 4, 0 \\le y \\le 4$).\n"
                 "Lượng sản phẩm I thu được: $30x + 20y$ (kg). Yêu cầu ít nhất 90 kg: $30x + 20y \\ge 90 \\iff 3x + 2y \\ge 9$.\n"
                 "Lượng sản phẩm II thu được: $30x + 40y$ (kg). Yêu cầu ít nhất 120 kg: $30x + 40y \\ge 120 \\iff 3x + 4y \\ge 12$.\n"
                 "Tổng chi phí mua nguyên liệu là $C(x, y) = 4x + 5y$ (triệu đồng).\n"
                 "Miền nghiệm của hệ ràng buộc là đa giác giới hạn bởi các đỉnh:\n"
                 "- Điểm $M(1/3; 4)$ (giao điểm của $y = 4$ và $3x + 2y = 9$): $C(1/3, 4) = 4(1/3) + 5(4) = \\dfrac{64}{3} \\approx 21.33$ triệu đồng.\n"
                 "- Điểm $N(2; 1.5)$ (giao điểm của $3x + 2y = 9$ và $3x + 4y = 12$):\n"
                 "  $C(2; 1.5) = 4(2) + 5(1.5) = 8 + 7.5 = 15.5$ triệu đồng.\n"
                 "- Điểm $P(4; 0)$ (giao điểm của $y = 0$ và $x = 4$, thỏa mãn $3(4) + 4(0) = 12 \\ge 12$ và $3(4) + 2(0) = 12 \\ge 9$):\n"
                 "  $C(4, 0) = 4(4) + 5(0) = 16$ triệu đồng.\n"
                 "- Điểm $Q(4; 4)$: $C(4, 4) = 4(4) + 5(4) = 36$ triệu đồng.\n"
                 "So sánh các chi phí tại các đỉnh: $15.5 < 16 < 21.33 < 36$.\n"
                 "Vậy phương án có chi phí nhỏ nhất là mua **2 tấn nguyên liệu K** và **1.5 tấn nguyên liệu P**, chi phí nhỏ nhất là **15.5 triệu đồng** (15 500 000 đồng).")

    add_sol_item("2 (Phần IV)", "Từ một tấm bìa hình tròn, cắt ra tam giác $ABC$ có $AB = 5\\text{ cm}, AC = 7\\text{ cm}$ và $\\widehat{B} = 60^\\circ$. Tính $BC$ và bán kính $R$ của hình tròn *(làm tròn $R$ đến hàng phần mười)*.",
                 "1) **Tính độ dài cạnh $BC$:**\n"
                 "Áp dụng định lí cosin trong tam giác $ABC$:\n"
                 "$AC^2 = AB^2 + BC^2 - 2AB \\cdot BC \\cos B$\n"
                 "$\\iff 7^2 = 5^2 + BC^2 - 2 \\cdot 5 \\cdot BC \\cdot \\cos 60^\\circ$\n"
                 "$\\iff 49 = 25 + BC^2 - 5BC \\iff BC^2 - 5BC - 24 = 0$.\n"
                 "Giải phương trình bậc hai: $(BC - 8)(BC + 3) = 0$.\n"
                 "Vì độ dài đoạn thẳng $BC > 0$, ta nhận $BC = 8\\text{ cm}$.\n\n"
                 "2) **Tính bán kính $R$ của hình tròn:**\n"
                 "Hình tròn bao quanh tam giác $ABC$ chính là đường tròn ngoại tiếp tam giác $ABC$. Áp dụng định lí sin:\n"
                 "$\\dfrac{AC}{\\sin B} = 2R \\implies R = \\dfrac{AC}{2\\sin B} = \\dfrac{7}{2\\sin 60^\\circ} = \\dfrac{7}{2 \\cdot \\frac{\\sqrt{3}}{2}} = \\dfrac{7}{\\sqrt{3}} = \\dfrac{7\\sqrt{3}}{3} \\approx 4.0415\\text{ cm}$.\n"
                 "Làm tròn đến hàng phần mười, ta được $R \\approx 4.0\\text{ cm}$.\n"
                 "**Kết luận:** $BC = 8\\text{ cm}$ và bán kính $R \\approx 4.0\\text{ cm}$.",
                 fig_path=r"C:\code\mathplus\lop10\figures_de294\p4_c2_circle.png", fig_w=Inches(2.5))

    add_sol_item("3 (Phần IV)", "Người ta lắp một thiết bị chiếu sáng sát tường, cao 5 m so với mặt đất. Góc chiếu sáng là $30^\\circ$ và dải sáng bắt đầu tại điểm cách chân tường 3 m. Độ dài vùng được chiếu sáng trên mặt đất bằng bao nhiêu mét? *(Làm tròn đến hàng phần mười)*.",
                 "Mô hình hóa hình học bài toán:\n"
                 "- Bức tường $AB$ vuông góc với mặt đất $DB$ tại $B$, với chiều cao thiết bị $AB = 5\\text{ m}$.\n"
                 "- Điểm bắt đầu dải sáng là $C$ trên mặt đất, cách chân tường $B$ một đoạn $BC = 3\\text{ m}$.\n"
                 "- Vùng được chiếu sáng trên mặt đất là đoạn $CD$, với điểm $D$ là điểm xa nhất dải sáng chiếu tới.\n"
                 "- Góc chiếu sáng của thiết bị là $\\widehat{CAD} = 30^\\circ$.\n\n"
                 "Xét tam giác vuông $ABC$ vuông tại $B$:\n"
                 "$\\tan(\\widehat{BAC}) = \\dfrac{BC}{AB} = \\dfrac{3}{5} = 0.6 \\implies \\widehat{BAC} = \\arctan(0.6) \\approx 30.9638^\\circ$.\n"
                 "Góc chiếu toàn phần từ tường tới mép ngoài của vùng sáng là:\n"
                 "$\\widehat{BAD} = \\widehat{BAC} + \\widehat{CAD} \\approx 30.9638^\\circ + 30^\\circ = 60.9638^\\circ$.\n"
                 "Xét tam giác vuông $ABD$ vuông tại $B$:\n"
                 "$BD = AB \\cdot \\tan(\\widehat{BAD}) = 5 \\cdot \\tan(60.9638^\\circ) \\approx 5 \\cdot 1.80125 \\approx 9.0063\\text{ m}$.\n"
                 "Độ dài vùng được chiếu sáng trên mặt đất là:\n"
                 "$CD = BD - BC = 9.0063 - 3 = 6.0063\\text{ m} \\approx 6.0\\text{ m}$.\n"
                 "**Kết luận:** Độ dài vùng được chiếu sáng trên mặt đất xấp xỉ **6.0 mét**.",
                 fig_path=r"C:\code\mathplus\lop10\figures_de294\p4_c3_light.png", fig_w=Inches(2.8))

    doc.save(output_path)
    print(f"Solution file created successfully at: {output_path}")

if __name__ == "__main__":
    out_sol = r"C:\code\mathplus\lop10\De_Kiem_Tra_Giua_HK1_Toan_10_THPT_Tran_Nhan_Tong_MaDe_294_DapAn_ChiTiet.docx"
    build_solution_doc(out_sol)
